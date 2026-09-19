# -*- coding: utf-8 -*-
"""Manual Maven and Website Edit Checklist — v1.2 -> v1.3.

The Maven page has now been independently checked by the owner, so v1.3 stops
hedging about it. Sections A and B record that the instructor title and bio are
already aligned and must not be touched. Section C replaces all three learning
outcomes against the verified live wording rather than fixing outcome 3 alone.
Section D's recommended copy no longer calls a group session private.

Section F gains the four site conflicts, each quoted from the file and line it
actually sits on in origin/main. maven.com and temidayoafonja.com remain blocked
by this environment's egress proxy, so the Maven values below are the owner's
verified reading rather than anything this pass read from a live page. The site
lines are read from the repository, which is the source the live site deploys
from.
"""
import copy, hashlib, os, shutil
import docx

SRC = "sept23-v206-assets/Manual_Maven_and_Website_Edit_Checklist_v1.2.docx"
OUT = "scratchpad/sept23/out"
DST = f"{OUT}/Manual_Maven_and_Website_Edit_Checklist_v1.3.docx"
STAMP = "Saturday, September 19, 2026 at 11:20 AM CT"
OLD = "Saturday, September 19, 2026 at 8:30 AM CT"

SUBS = [("v1.2", "v1.3"), (OLD, STAMP), ("deck v2.0.6", "deck v2.0.7")]

PARA = [
 (1, "NOTHING IN THIS CHECKLIST HAPPENS AUTOMATICALLY.",
  "NOTHING IN THIS CHECKLIST HAPPENS AUTOMATICALLY. The repository has no route "
  "to publish to Maven or to the live site. The Maven values in sections A to E "
  "are the owner's own verified reading of the live page, not anything this pass "
  "read: maven.com and temidayoafonja.com are both blocked by this environment's "
  "network egress. The website lines in section F are quoted from the repository "
  "with their file and line numbers, and were confirmed against origin/main."),
 (4, "ALREADY CURRENT — DO NOT OVERWRITE.",
  "ALREADY CURRENT — DO NOT OVERWRITE. Verified on the live page. The instructor "
  "title is aligned and must remain unchanged. The table below is kept only as "
  "the record of when it changed."),
 (8, "ALREADY CURRENT — DO NOT OVERWRITE.",
  "ALREADY CURRENT — DO NOT OVERWRITE. Verified on the live page. The bio is "
  "aligned and must remain unchanged. Do not paste the version below over it; it "
  "is kept only as the record."),
 (16, "Recognize the experienced professional before teaching anything.",
  "Recognize the experienced professional before teaching anything, and do not "
  "call the session private. Private SCORING is accurate and should be said. "
  "Calling the group session itself private is not. Use “live, evidence-based "
  "session” or “guided, evidence-based session”."),
 (18, "Suggested direction, to adapt rather than paste:",
  "Suggested direction, to adapt rather than paste: experienced professionals can "
  "be performing well, trusted with important work and carrying real "
  "responsibility while becoming less certain about what the work is building for "
  "the future. The difficult question is not only whether the job is good or bad. "
  "It is whether the work is still expanding your judgment, what would remain "
  "useful if the context changed, and what your recent evidence supports testing "
  "next. This live, evidence-based session uses your own last 90 days of work to "
  "help you read that more clearly. Your scoring stays private throughout."),
 (33, "The Private Capability Position Read is retired as an offer",
  "The Private Capability Position Read is retired as an offer and is still "
  "publicly visible on two pages. Removing it is the largest item in section F "
  "and cannot be done from this package. Do not substitute a Career Move Review "
  "call to action in its place: that route does not exist yet, which is also why "
  "deck v2.0.7 keeps Career Move Review off the participant-facing continuation "
  "slide."),
]

# Section C — all three outcomes, against the verified live wording.
OUTCOMES = [
 ("Live wording, verified", "Replace with"),
 ("Know whether your job is still building you",
  "OUTCOME 1 — See whether your current work is still building you. Supporting "
  "copy: use evidence from your last 90 days of actual work to distinguish "
  "continued formation from familiar performance."),
 ("Know what you can carry into another role",
  "OUTCOME 2 — See what may still count when the context changes. Supporting "
  "copy: examine what may travel beyond your current role or employer, where the "
  "evidence is strong, and what may still require testing or relearning."),
 ("Decide your next career move",
  "OUTCOME 3 — Identify which move is worth testing next. Supporting copy: "
  "evaluate seven move categories and leave with a Next-Move Note naming the "
  "direction your evidence supports testing next."),
 ("Why all three",
  "The session does not make the stay-or-leave decision for the participant and "
  "the listing must not promise that it does. Outcome 2 must not promise "
  "destination-specific portability when no destination has been named. This is "
  "the same boundary the deck and the SOP hold."),
 ("Do not change",
  "Date, time, duration and instructor positioning. Wednesday, September 23, "
  "2026, 6:00 PM CT, 60 minutes, free."),
]

# Section F — the four verified conflicts, with file and line references.
WEBSITE = [
 ("Item", "Change"),
 ("A. Private Read is retired but publicly visible — LARGEST ITEM",
  "It must come out of the current public offer architecture. fieldkit.html:337 "
  "carries “The Private Capability Position Read” with an “Enquire about the "
  "Private Read” call to action at line 349 pointing to work.html#get-in-touch, "
  "and line 348 still frames it as the choice beside the Field Kit. "
  "for-professionals.html:255 carries a Private Capability Position Read block "
  "with a mailto call to action. content/site-source-of-truth.json describes it "
  "as a current offer. DO NOT substitute a Career Move Review call to action "
  "until that route is operational and verified: removing the retired offer and "
  "adding a new one are two separate decisions."),
 ("B. /fieldkit defines Optionality partly as visibility",
  "fieldkit.html:230 reads “Optionality, your market portability. Whether that "
  "capability is yours to carry, and whether the people who could use it have "
  "seen it.” Flag for revision. Visibility is not evidence and should not define "
  "portability. The deck now reads Optionality as usefulness beyond the current "
  "context plus legibility through evidence, and the site should follow."),
 ("C. /for-professionals combines two different offers",
  "for-professionals.html:222 reads “Keep the Proof: The 10-Minute Career "
  "Evidence Starter”, linking to /career-evidence-starter. That is one label over "
  "two products. Current architecture: Career Evidence Starter is FREE, Keep the "
  "Proof is $49. Flag for correction so the free entry artifact and the paid "
  "product are named separately."),
 ("D. /for-professionals promises a decision",
  "for-professionals.html:217 reads “See whether your role is still building your "
  "capability, identify what would travel with you, and decide your next move.” "
  "Replace with language consistent with the listing: see what may still count "
  "when the context changes, and identify which move is worth testing next."),
 ("Career Move Review — does not exist yet",
  "No page, no route, no copy, on this branch or on origin/main. Until a verified "
  "individual-professional route exists it stays off the participant-facing "
  "continuation slide and is not sold from the stage. Do not route it to the "
  "enterprise Advisory page and do not reuse a retired Private Read link."),
 ("Field Kit route",
  "temidayoafonja.com/fieldkit is the participant-facing route on the deck face, "
  "the QR target and the prebuilt chat message. Confirm the page is live, "
  "resolves without a redirect chain, and shows the current $150 price."),
 ("Preserve unchanged",
  "September 23, free, 60 minutes, current live time, Field Kit $150, and every "
  "individual-professional call to action reaching an individual path rather than "
  "the enterprise Advisory page."),
]


def set_para(para, text):
    runs = para.runs
    keep = next((r for r in runs if r.text.strip()), runs[0])
    keep.text = text
    for r in runs:
        if r._r is not keep._r:
            r._r.getparent().remove(r._r)


def set_cell(cell, text):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    set_para(p, text)


def fill(table, rows):
    template = copy.deepcopy(table.rows[1]._element)
    while len(table.rows) > 1:
        table._element.remove(table.rows[-1]._element)
    for ci, v in enumerate(rows[0]):
        set_cell(table.rows[0].cells[ci], v)
    for row in rows[1:]:
        table._element.append(copy.deepcopy(template))
        for ci, v in enumerate(row):
            set_cell(table.rows[-1].cells[ci], v)


def build():
    os.makedirs(OUT, exist_ok=True)
    shutil.copyfile(SRC, DST)
    d = docx.Document(DST)

    def sub_runs(runs):
        for r in runs:
            t = r.text
            for a, b in SUBS:
                t = t.replace(a, b)
            if t != r.text:
                r.text = t

    for p in d.paragraphs:
        sub_runs(p.runs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    sub_runs(p.runs)

    for idx, expect, new in PARA:
        para = d.paragraphs[idx]
        assert expect in para.text, \
            f"P{idx}: expected {expect[:40]!r}, found {para.text[:70]!r}"
        set_para(para, new)

    assert "Decide your next career move" in d.tables[2].rows[1].cells[1].text
    fill(d.tables[2], OUTCOMES)
    assert "Private Read" in d.tables[3].rows[1].cells[0].text
    fill(d.tables[3], WEBSITE)

    d.save(DST)
    return DST


if __name__ == "__main__":
    path = build()
    print("built", os.path.basename(path))
    print("  sha256", hashlib.sha256(open(path, "rb").read()).hexdigest())
