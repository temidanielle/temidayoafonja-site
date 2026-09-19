# -*- coding: utf-8 -*-
"""Stay or Leave? 60-minute flagship — v2.0.6 -> v2.0.7 CANDIDATE.

Two changes, both narrow.

  SLIDE 22   The four-question portability frame moves out of a compressed line
             under the subtitle and onto its own restrained strip: the same
             light panel and gold accent the three prompt blocks already use,
             at body size rather than a small bold aside. It is still one line,
             still one occurrence, and the writing space below it is untouched.

  SLIDE 5    The presenter note gains one clarification: the Optionality score
             is a composite reading of portability conditions and signals, not
             a prediction that another employer will value, hire, promote or
             pay anyone. Said once, in the note, not on the participant-facing
             slide.

Nothing else is reopened. The opening, the Density and Optionality sequencing,
the slide 18 simplification, the protected twelve-minute block, the seven move
categories, the two-route continuation, the slide 25 close, the workbook and the
recording architecture are all left exactly as v2.0.6 approved them.
"""
import copy, hashlib, os, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC = ("sept23-v206-assets/PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_"
       "Assessment_60MIN_v2.0.6_CANDIDATE.pptx")
OUT = "scratchpad/sept23/out"
DST = (f"{OUT}/PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_"
       "60MIN_v2.0.7_CANDIDATE.pptx")

NAVY = RGBColor(0x0F, 0x23, 0x47)
FRAME = ("What travels?      What does not?      What can I prove?      "
         "What must I relearn?")

# Slide 22 geometry, in inches. The title box loses the slack it was not using,
# which is where the strip comes from. The three prompt blocks at 2.02, 2.84 and
# 3.66 and the closing band at 4.54 are not moved: the writing space is the
# point of the slide.
S22 = {
    1: (0.66, 0.62, 8.76, 0.64),     # title
    2: (1.32, 0.62, 8.76, 0.24),     # subtitle
}
# The strip sits 0.10in clear of the first prompt block, which is exactly the
# gap the three blocks already keep between themselves. Flush against it, the
# frame read as a heading for MY CURRENT READ rather than for the whole note.
STRIP_PANEL = (1.62, 0.62, 8.76, 0.30)
STRIP_ACCENT = (1.62, 0.62, 0.04, 0.30)
STRIP_TEXT = (1.66, 0.86, 8.28, 0.24)

OPTIONALITY_NOTE = (
 "\n\nWHAT THE OPTIONALITY SCORE IS, AND SAY THIS ONCE IF IT COMES UP: the six "
 "statements are a composite reading of portability conditions and signals. They "
 "are not a prediction. Nothing on this axis forecasts that another employer will "
 "value, hire, promote or pay anyone. If someone treats the number as a market "
 "verdict, correct it plainly: it reads the conditions you are working in, not a "
 "decision nobody has made yet.\n"
 "INTERNAL, NOT FOR THE ROOM: statements 7, 10 and 12 carry anticipated employer "
 "valuation, visibility and a one-year forecast, which sit less cleanly inside "
 "that reading than the other three. They are recorded in the change log as a "
 "post-September-23 instrument-review item. Do NOT reword them in delivery: the "
 "statements are locked and the workbook is built on them.")


def set_text(shape, text):
    tf = shape.text_frame
    first = tf.paragraphs[0]
    keep = next((r for r in first.runs if r.text.strip()), first.runs[0])
    for p in tf.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    for r in list(first.runs):
        if r._r is not keep._r:
            r._r.getparent().remove(r._r)
    keep.text = text
    return keep


def place(shape, geom):
    top, left, w, h = geom
    shape.top, shape.left, shape.width, shape.height = (
        Inches(top), Inches(left), Inches(w), Inches(h))


def clone(slide, source_shape, geom):
    """Copy an existing shape so the strip inherits the deck's own styling
    rather than being drawn from scratch with guessed colours."""
    el = copy.deepcopy(source_shape._element)
    slide.shapes._spTree.append(el)
    shp = slide.shapes[-1]
    place(shp, geom)
    return shp


def build():
    os.makedirs(OUT, exist_ok=True)
    shutil.copyfile(SRC, DST)
    p = Presentation(DST)
    sl = list(p.slides)
    assert len(sl) == 33, "slide count moved"

    # ── slide 22 ────────────────────────────────────────────────────────────
    s22 = sl[21]
    shp = list(s22.shapes)
    assert "The Next-Move Note" in shp[1].text_frame.text, "slide 22 title moved"
    assert "What travels?" in shp[2].text_frame.text, \
        "slide 22 subtitle does not carry the v2.0.6 frame"
    assert "MY CURRENT READ" in shp[5].text_frame.text, "prompt blocks moved"

    for j, geom in S22.items():
        place(shp[j], geom)
    # the subtitle goes back to one line; the frame gets its own strip
    set_text(shp[2], "Three lines. Yours, private, and enough to act on.")

    clone(s22, shp[3], STRIP_PANEL)      # the light panel the prompts use
    clone(s22, shp[4], STRIP_ACCENT)     # its gold accent bar
    run = set_text(clone(s22, shp[6], STRIP_TEXT), FRAME)
    run.font.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = NAVY

    # ── slide 5 ─────────────────────────────────────────────────────────────
    n5 = sl[4].notes_slide.notes_text_frame
    assert "Optionality" in "\n".join(s.text_frame.text for s in sl[4].shapes
                                      if s.has_text_frame), "slide 5 is not Optionality"
    assert "composite reading" not in n5.text, "slide 5 already carries the note"
    n5.text = n5.text + OPTIONALITY_NOTE

    p.save(DST)

    # ── prove the pass was narrow ───────────────────────────────────────────
    before, after = Presentation(SRC), Presentation(DST)
    bf = ["\n".join(s.text_frame.text for s in x.shapes if s.has_text_frame)
          for x in before.slides]
    af = ["\n".join(s.text_frame.text for s in x.shapes if s.has_text_frame)
          for x in after.slides]
    bn = [x.notes_slide.notes_text_frame.text if x.has_notes_slide else ""
          for x in before.slides]
    an = [x.notes_slide.notes_text_frame.text if x.has_notes_slide else ""
          for x in after.slides]
    faces = [i + 1 for i in range(33) if bf[i] != af[i]]
    notes = [i + 1 for i in range(33) if bn[i] != an[i]]
    assert faces == [22], f"a face changed outside slide 22: {faces}"
    assert notes == [5], f"a note changed outside slide 5: {notes}"
    assert an[4] == bn[4] + OPTIONALITY_NOTE, "slide 5's note gained something else"

    print("built", os.path.basename(DST))
    print(f"  faces changed: {faces}   notes changed: {notes}")
    print("  sha256", hashlib.sha256(open(DST, "rb").read()).hexdigest())
    return DST


if __name__ == "__main__":
    build()
