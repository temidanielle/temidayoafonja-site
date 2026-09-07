# -*- coding: utf-8 -*-
"""Document generation for the new Videos 4 and 5 recording packages."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x11, 0x23, 0x45)
GOLD = RGBColor(0x8A, 0x6D, 0x1E)
DIM  = RGBColor(0x5A, 0x6B, 0x82)
RED  = RGBColor(0x9B, 0x2C, 0x10)


def base_doc():
    d = Document()
    s = d.sections[0]
    s.top_margin = s.bottom_margin = Inches(0.8)
    s.left_margin = s.right_margin = Inches(0.9)
    n = d.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(11)
    n._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    n.paragraph_format.space_after = Pt(8)
    n.paragraph_format.line_spacing = 1.18
    return d


def para(d, text, size=11, bold=False, color=None, before=0, after=8,
         align=None, italic=False, keep=False):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER,
                       "r": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if keep:
        pr = p._p.get_or_add_pPr()
        for tag in ("keepNext", "keepLines"):
            e = OxmlElement("w:" + tag); e.set(qn("w:val"), "1"); pr.append(e)
    return p


def title_block(d, eyebrow, title, sub=None):
    para(d, eyebrow.upper(), size=9, bold=True, color=GOLD, after=4)
    para(d, title, size=21, bold=True, color=NAVY, after=4)
    if sub:
        para(d, sub, size=10, color=DIM, after=14)
    rule(d)


def rule(d, after=12):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    pr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "6")
    bt.set(qn("w:color"), "8A6D1E"); b.append(bt); pr.append(b)


def h(d, text, size=13):
    para(d, text.upper(), size=size, bold=True, color=NAVY, before=16, after=6,
         keep=True)


def kv(d, k, v):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(k + "  ")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = GOLD
    r2 = p.add_run(v); r2.font.size = Pt(10.5)
    return p


def callout(d, text, color=RED):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.15)
    pr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); lt = OxmlElement("w:left")
    lt.set(qn("w:val"), "single"); lt.set(qn("w:sz"), "18")
    lt.set(qn("w:color"), "9B2C10"); lt.set(qn("w:space"), "8")
    b.append(lt); pr.append(b)
    for tag in ("keepNext", "keepLines"):
        e = OxmlElement("w:" + tag); e.set(qn("w:val"), "1"); pr.append(e)
    r = p.add_run(text); r.font.size = Pt(10.5); r.bold = True
    r.font.color.rgb = color
    return p
