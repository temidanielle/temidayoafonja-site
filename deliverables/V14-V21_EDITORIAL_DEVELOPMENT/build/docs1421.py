# -*- coding: utf-8 -*-
"""Document helpers for the V14 to V21 editorial development deliverables.

Built on the shared house helpers in new-videos-4-5/build/docs.py. Table and
bullet helpers are added here rather than in the shared file so no other
batch's documents can change.
"""
import os, sys
DELIV = "/home/user/temidayoafonja-site/deliverables/"
sys.path.insert(0, DELIV + "new-videos-4-5/build")

from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docs import (base_doc, para, title_block, rule, h, kv, callout,
                  NAVY, GOLD, DIM, RED)

CREAMHEX = "F3EFE6"
LINEHEX = "C9BFA8"


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


# w:tblPr children are order-sensitive. w:tblBorders must sit before
# w:tblLayout, w:tblCellMar and w:tblLook, and appending it at the end
# produces a file LibreOffice refuses to open, so insert it by position.
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
    "tblDescription",
]


def _insert_ordered(parent, element, order):
    """Insert element into parent at its schema position."""
    name = element.tag.split("}")[-1]
    idx = order.index(name)
    for child in parent:
        cname = child.tag.split("}")[-1]
        if cname in order and order.index(cname) > idx:
            child.addprevious(element)
            return
    parent.append(element)


def _borders(table):
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), LINEHEX)
        borders.append(e)
    _insert_ordered(pr, borders, TBLPR_ORDER)


def _cell_text(cell, text, size=9, bold=False, color=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color


def table(d, header, rows, widths=None, size=9, head_size=8.5):
    """A bordered table. widths are inches and must sum to the text width."""
    t = d.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _borders(t)
    for i, htext in enumerate(header):
        c = t.rows[0].cells[i]
        _shade(c, CREAMHEX)
        _cell_text(c, htext.upper(), size=head_size, bold=True, color=NAVY)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            _cell_text(cells[i], val, size=size)
    if widths:
        for r_ in t.rows:
            for i, w in enumerate(widths):
                r_.cells[i].width = Inches(w)
    # keep header rows visible when a table breaks across pages
    trPr = t.rows[0]._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)
    d.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def bullets(d, items, size=10.5, indent=0.22, after=4):
    for it in items:
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(indent + 0.16)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.16
        r = p.add_run("•   " + it)
        r.font.size = Pt(size)


def sub(d, text, size=10.5):
    """A small gold label above a block."""
    return para(d, text.upper(), size=8.5, bold=True, color=GOLD,
                before=10, after=3, keep=True)


def field(d, label, text, size=10.5):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.16
    r = p.add_run(label + "  ")
    r.bold = True
    r.font.size = Pt(size - 0.5)
    r.font.color.rgb = NAVY
    r2 = p.add_run(text)
    r2.font.size = Pt(size)
    return p


def page_break(d):
    from docx.enum.text import WD_BREAK
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def footer_note(d, text):
    sec = d.sections[0]
    f = sec.footer.paragraphs[0]
    f.text = ""
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run(text)
    r.font.size = Pt(8)
    r.font.color.rgb = DIM
