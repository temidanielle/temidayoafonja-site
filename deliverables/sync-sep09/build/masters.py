# -*- coding: utf-8 -*-
"""The four September 9 locked Recording Masters, parsed.

These files are the spoken source of truth. Nothing here rewrites a line of
them: the module only reads them, so that every sentence trigger, run of show
entry and reading copy in the package is generated from the locked text rather
than transcribed by hand.
"""
import os, re, json, hashlib
from docx import Document

HANDOFF = ("/tmp/claude-0/-home-user-temidayoafonja-site/"
           "f121668d-e262-5eb8-9b22-0eaa1006a361/scratchpad/handoff")
MASTERS_DIR = os.path.join(HANDOFF, "Recording_Masters")
MANIFEST = json.load(open(os.path.join(HANDOFF, "Source_Lock_Manifest.json")))

CANON = {r["video"]: r["canonical_filename"] for r in MANIFEST["recording_sources"]}
SRC = {r["video"]: r for r in MANIFEST["recording_sources"]}

# Bracketed directions are production instructions, not spoken dialogue.
DIRECTION = re.compile(r"^\[.*\]$")
# All-caps section markers in V6 and V7 are production markers, not speech.
CAPS_MARKER = re.compile(r"^[A-Z0-9][A-Z0-9 ,:?+/'&.-]{2,59}$")
END_MARK = re.compile(r"^END\b")


def path(n):
    return os.path.join(MASTERS_DIR, CANON[n])


def verify(n):
    """The canonical file must still hash to what the lock manifest recorded."""
    h = hashlib.sha256(open(path(n), "rb").read()).hexdigest()
    return h, h == SRC[n]["canonical_sha256"]


def _rows(n):
    """(style, text) for every non-empty paragraph of the master."""
    d = Document(path(n))
    return [(p.style.name, p.text.strip()) for p in d.paragraphs if p.text.strip()]


def tables(n):
    out = []
    for t in Document(path(n)).tables:
        out.append([[c.text.strip() for c in r.cells] for r in t.rows])
    return out


def header(n):
    """The master's own cover table, as a dict."""
    out = {}
    for tb in tables(n):
        for row in tb:
            if len(row) == 2 and row[0] and row[1] and len(row[0]) < 40:
                out[row[0].strip().upper()] = row[1].strip()
    return out


def visual_map(n):
    """The master's own visual / motion-graphic map rows."""
    for tb in tables(n):
        if tb and [c.strip().lower() for c in tb[0][:2]] == ["#", "visual job"]:
            return [dict(num=r[0].strip(), job=r[1].strip(),
                         onscreen=r[2].strip(), treatment=r[3].strip())
                    for r in tb[1:]]
    return []


def sections(n):
    """[(section title, [(kind, text), ...]), ...] for the Recording Script only.

    kind is "speech" for spoken copy, "direction" for a bracketed delivery
    direction, and "end" for the closing production note. Section titles come
    from the master's own headings, so nothing is invented or reordered.
    """
    rows = _rows(n)
    start = next(i for i, (s, t) in enumerate(rows)
                 if s == "Heading 1" and t == "Recording Script") + 1
    stop = len(rows)
    for i in range(start, len(rows)):
        if rows[i][0] in ("Heading 1", "Heading 2") and \
           rows[i][1].startswith(("Major Visual", "Recording Close")):
            stop = i
            break
    out, cur = [], None

    def open_section(title):
        nonlocal cur
        cur = (title, [])
        out.append(cur)

    open_section("")
    for style, text in rows[start:stop]:
        if style in ("Heading 1", "Heading 2"):
            open_section(text)
        elif n in (6, 7) and CAPS_MARKER.match(text) and not text.endswith("."):
            open_section(text)
        elif DIRECTION.match(text):
            cur[1].append(("direction", text))
        elif END_MARK.match(text):
            cur[1].append(("end", text))
        else:
            cur[1].append(("speech", text))
    return [(t, body) for t, body in out if body]


def speech(n):
    return [t for _, body in sections(n) for k, t in body if k == "speech"]


def directions(n):
    return [(title, t) for title, body in sections(n)
            for k, t in body if k == "direction"]


def word_count(n):
    return sum(len(t.split()) for t in speech(n))


def contains(n, sentence):
    """Is this exact sentence in the locked spoken script?

    Curly and straight quotes are normalized, because package documents are
    written with straight quotes while the master uses typographic ones.
    """
    def norm(s):
        return (s.replace("’", "'").replace("‘", "'")
                 .replace("“", '"').replace("”", '"'))
    body = norm("\n".join(speech(n)))
    return norm(sentence).strip() in body


if __name__ == "__main__":
    for n in (4, 5, 6, 7):
        h, ok = verify(n)
        secs = sections(n)
        print("V%d  %s" % (n, CANON[n]))
        print("    sha256 matches lock manifest: %s" % ok)
        print("    sections=%d  speech=%d  directions=%d  words=%d "
              "(manifest says approx %d)"
              % (len(secs), len(speech(n)), len(directions(n)), word_count(n),
                 SRC[n]["approx_spoken_word_count"]))
        print("    visual map rows: %d" % len(visual_map(n)))
        for t, body in secs:
            print("      %-46s %d" % (("(open)" if not t else t)[:46], len(body)))
        for title, d in directions(n):
            print("      DIRECTION in %-28s %s" % (title[:28], d))
        print()
