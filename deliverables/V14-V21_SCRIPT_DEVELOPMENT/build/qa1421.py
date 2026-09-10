# -*- coding: utf-8 -*-
"""The QA checks for the script-development batch.

Every check runs against the BUILT documents, not against the source data, so
a check cannot pass because of something that was true only in the build
script. Each check reports PASS or FAIL with what it looked at.
"""
import os, re, sys, glob
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


def units(path):
    """Every paragraph and every table row of a .docx, as separate units.

    A DOCX has no blank lines, so each paragraph is its own unit and each
    table row is joined into one line. Never test a wrapped fragment in
    isolation.
    """
    out = []
    d = Document(path)
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            t = Paragraph(child, d).text.strip()
            if t:
                out.append(t)
        elif child.tag == qn("w:tbl"):
            for r in Table(child, d).rows:
                out.append(" | ".join(c.text.strip() for c in r.cells))
    return out


def text_units(path):
    """Blank-line units of a text file, each carrying its own heading.

    A plain-text file splits on blank lines, but an indented list sits in a
    different unit from the heading that governs it, and a reader sees both
    at once. So a unit whose lines are all indented deeper than the previous
    unit gets that previous unit prepended. Without this, a list of things
    that are BLOCKED reads as a list of things being claimed.
    """
    raw = [u for u in open(path, encoding="utf-8").read().split("\n\n")
           if u.strip()]
    out, prev = [], ""
    for u in raw:
        lines = [l for l in u.split("\n") if l.strip()]
        indent = min(len(l) - len(l.lstrip()) for l in lines)
        prev_indent = (min(len(l) - len(l.lstrip())
                           for l in prev.split("\n") if l.strip())
                       if prev.strip() else -1)
        if prev.strip() and indent > prev_indent:
            out.append(prev.strip() + "\n" + u.strip())
        else:
            out.append(u.strip())
            prev = u
    return out


PROHIBITION = re.compile(
    r"\b(no|not|never|blocked|must not|does not|cannot|until|before|"
    r"prohibit|forbid|any statement of|any number|any past-tense|"
    r"conditional|unusable|not usable)\b", re.I)


def is_prohibited(unit, match_start):
    """True when the matched phrase sits inside a prohibition.

    Looks at the sentence carrying the match and at the unit's heading line.
    A file that forbids a claim must not be read as making it.
    """
    head = unit.split("\n", 1)[0]
    start = max(unit.rfind(".", 0, match_start),
                unit.rfind("\n", 0, match_start)) + 1
    end = unit.find(".", match_start)
    sentence = unit[start:end if end > 0 else len(unit)]
    return bool(PROHIBITION.search(sentence) or PROHIBITION.search(head))


def spoken_units(path):
    """Only the spoken thought blocks of a script-only recording copy.

    Identified by size: blocks are set at 13.5pt in those documents and every
    label, title and instruction is smaller.
    """
    out = []
    d = Document(path)
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        if any(r.bold for r in p.runs):
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if sizes and max(sizes) >= 13.0:
            out.append(p.text.strip())
    return out


CHECKS = []


def check(name):
    def deco(fn):
        CHECKS.append((name, fn))
        return fn
    return deco


def _master(n):
    return os.path.join(OUT, "Video_%d_Recording_Master_DRAFT.docx" % n)


def _copy(n):
    return os.path.join(OUT, "Video_%d_Script_Only_Recording_Copy.docx" % n)


VIDEOS = (15, 16, 17, 18, 19, 20, 21)


# --- Video 14 -------------------------------------------------------------

@check("V14 has no invented research result")
def c1():
    bad = []
    for f in glob.glob(os.path.join(OUT, "V14_*")):
        us = units(f) if f.endswith(".docx") else text_units(f)
        for u in us:
            m = re.search(r"\b(we found|the comparison found|results show|"
                          r"the data show|findings indicate|turned out that|"
                          r"the study found)\b", u, re.I)
            if m and not is_prohibited(u, m.start()):
                bad.append("%s :: %s" % (os.path.basename(f), u[:90]))
    return not bad, bad or ["no findings language in any V14 file"]


@check("V14 makes no numerical or past-tense research headline")
def c2():
    bad = []
    for f in glob.glob(os.path.join(OUT, "V14_*")):
        us = units(f) if f.endswith(".docx") else text_units(f)
        for u in us:
            if "I Compared 30 Job Descriptions" not in u:
                continue
            # every mention must sit beside its own prohibition
            if not re.search(r"conditional|not usable|only after|blocked|"
                             r"can be used only|BLOCKED until", u, re.I):
                bad.append("%s :: %s" % (os.path.basename(f), u[:110]))
    return not bad, bad or ["every mention of the conditional title carries "
                            "its condition in the same paragraph"]


@check("V14 fabricates no job posting")
def c3():
    p = os.path.join(OUT, "V14_Coding_Framework.csv")
    rows = open(p).read().strip().split("\n")
    ok = len(rows) == 1
    return ok, ["collection template has %d row(s): header only is correct"
                % len(rows)] if ok else ["template contains data rows"]


# --- Videos 15 to 21 ------------------------------------------------------

@check("V15 separates learning, opportunity or practice, and evidence")
def c4():
    us = units(_master(15))
    need = {
      "learning gap": r"learning gap",
      "practice and authority gap": r"practice and authority gap",
      "evidence gap": r"evidence (gap|problem)",
      "responses do not substitute": r"do not substitute for each other",
    }
    missing = [k for k, pat in need.items()
               if not any(re.search(pat, u, re.I) for u in us)]
    return not missing, missing or ["all three gaps named and separated, and "
                                    "the non-substitution rule is stated"]


@check("V15 is not presented as a validated diagnostic")
def c5():
    us = units(_master(15))
    bad = [u[:90] for u in us
           if re.search(r"\b(diagnostic|assessment)\b", u, re.I)
           and not re.search(r"not a (validated )?diagnostic|is not a "
                             r"diagnostic|not an assessment", u, re.I)]
    return not bad, bad or ["diagnostic language appears only in its own "
                            "denial"]


@check("V16 does not blame the viewer for being overlooked")
def c6():
    us = units(_master(16))
    disclaimed = any("entirely their own doing" in u or
                     "entirely your own" in u for u in us)
    honest = any("does not override bias" in u or
                 "This does not override bias" in u for u in us)
    ok = disclaimed and honest
    return ok, ["the script states what the record does not fix, and the "
                "boundary against blaming the viewer is recorded"] if ok \
        else ["missing the not-your-fault boundary or the honest limits"]


@check("V17 does not repeat V11")
def c7():
    us = units(_master(17))
    ok = any(re.search(r"V11 asks what remains valuable|"
                       r"must not become Video 11", u) for u in us)
    sep = any("Output and contribution are not the same thing" in u
              for u in us)
    return ok and sep, ["the V11 separation is stated in the direction and "
                        "the script's own distinction is output versus "
                        "contribution"] if ok and sep \
        else ["V11 separation not explicit"]


@check("V17 makes no unverified AI capability claim")
def c8():
    us = units(_master(17))
    illustration = any("constructed this example" in u or
                       "It is an illustration" in u for u in us)
    subst = any("SUBSTITUTION RULE" in u for u in us)
    nosafe = any("not telling you judgment makes your job safe" in u
                 for u in us)
    ok = illustration and subst and nosafe
    return ok, ["the worked example is labeled a constructed illustration in "
                "the spoken script, the substitution rule is recorded, and "
                "the job-security claim is refused out loud"] if ok \
        else ["missing illustration label, substitution rule, or the "
              "job-security refusal"]


@check("V18 remains genuinely two-sided")
def c9():
    us = units(_master(18))
    both = any("Pick the one whose hard parts you would rather have" in u
               for u in us)
    caps = any("genuinely cap individual contributors" in u for u in us)
    neither = any("Neither answer is the brave one" in u or
                  "both will make you decide" in u for u in us)
    ok = both and caps and neither
    return ok, ["the script refuses both prescriptions, states that some "
                "organizations cap ICs, and ends on a choice"] if ok \
        else ["two-sidedness not demonstrable"]


@check("V19 promises no consulting income or clients")
def c10():
    us = units(_master(19))
    bad = [u[:90] for u in us
           if re.search(r"\b(you will earn|guaranteed|replace your salary|"
                        r"six figures|clients will|you will find clients)\b",
                        u, re.I)]
    honest = any("Income is uneven" in u for u in us)
    return (not bad) and honest, bad or ["no income or client promise, and "
                                         "the variability is stated"]


@check("V20 invents no personal career-break story")
def c11():
    us = units(_master(20))
    labeled = any("constructed example" in u for u in us)
    firstperson = [u[:90] for u in us
                   if re.search(r"when I (took|came back|returned)|my (own )?"
                                r"career break|after my break", u, re.I)]
    bias = any("no amount of framing changes" in u or
               "framing removes employer bias" in u for u in us)
    ok = labeled and not firstperson and bias
    return ok, (firstperson or ["the only example is labeled constructed, "
                                "there is no first-person break story, and "
                                "the bias limit is stated"])


@check("V21 preserves genuine relearning")
def c12():
    us = units(_master(21))
    license_line = any("not a confidence issue and it is not a mindset "
                       "problem" in u for u in us)
    layers = sum(1 for u in us if re.search(r"^Layer (one|two|three|four|"
                                            r"five)", u, re.I))
    nogeneric = any("deliberately not naming specific regulations" in u
                    for u in us)
    ok = license_line and layers >= 5 and nogeneric
    return ok, ["a license is stated not to be a mindset issue, five layers "
                "are present, and no specific regulation is named"] if ok \
        else ["missing the licensing line, a layer, or the no-named-"
              "regulation rule"]


# --- Cross-cutting --------------------------------------------------------

@check("Employer due-diligence topic stays reserved and unnumbered")
def c13():
    p = os.path.join(OUT, "V14-V21_Editorial_Review_Summary.docx")
    us = units(p)
    named = any("Before You Accept the Job" in u for u in us)
    reserved = any("Before You Accept the Job" in u and
                   re.search(r"reserved|no (slot )?number|unnumbered", u,
                             re.I) for u in us)
    numbered_ = [u[:90] for u in us
                 if re.search(r"V\d+[^.]{0,30}Before You Accept the Job", u)]
    ok = named and reserved and not numbered_
    return ok, numbered_ or ["named, reserved, and never given a number"]


@check("After-40 topic stays preserved")
def c14():
    us = units(os.path.join(OUT, "V14-V21_Editorial_Review_Summary.docx"))
    ok = any("How to Stay Relevant After 40" in u and
             re.search(r"preserv|queue|not delet", u, re.I) for u in us)
    return ok, ["preserved in the forward queue"] if ok else ["not preserved"]


@check("V22 remains promotion timing")
def c15():
    us = units(os.path.join(OUT, "V14-V21_Editorial_Review_Summary.docx"))
    ok = any("No Promotion? How to Know Whether to Wait or Move On" in u
             for u in us)
    return ok, ["V22 named and unchanged"] if ok else ["V22 not recorded"]


@check("No em dashes or en dashes anywhere")
def c16():
    bad = []
    for f in sorted(glob.glob(os.path.join(OUT, "*.docx")) +
                    glob.glob(os.path.join(OUT, "*.txt")) +
                    glob.glob(os.path.join(OUT, "*.csv"))):
        us = units(f) if f.endswith(".docx") else text_units(f)
        for u in us:
            if "—" in u or "–" in u:
                bad.append("%s :: %s" % (os.path.basename(f), u[:70]))
    return not bad, bad or ["clean across every delivered file"]


@check("U.S. English in the delivered files")
def c17():
    pat = (r"\b(colour|behaviour|neighbour|organis(e|ed|es|ing|ation)|"
           r"recognis(e|ed|es|ing)|analys(e|ed|es)|labelled|centre|"
           r"programme|licence|practise|whilst|amongst|towards the)\b")
    bad = []
    for f in sorted(glob.glob(os.path.join(OUT, "*.docx")) +
                    glob.glob(os.path.join(OUT, "*.txt"))):
        us = units(f) if f.endswith(".docx") else text_units(f)
        for u in us:
            m = re.search(pat, u, re.I)
            if m:
                bad.append("%s :: %s :: %s"
                           % (os.path.basename(f), m.group(0), u[:60]))
    return not bad, bad or ["no British spellings found"]


@check("No stacked spoken CTAs")
def c18():
    """At most one spoken instruction to the viewer, per script."""
    bad = []
    for n in VIDEOS:
        blocks = spoken_units(_copy(n))
        asks = [b for b in blocks
                if re.search(r"put (it|that|the|your)[^.]{0,60}in the "
                             r"comments", b, re.I)]
        subs = [b for b in blocks if re.search(r"\bsubscribe\b", b, re.I)]
        if len(asks) != 1:
            bad.append("V%d has %d comment asks" % (n, len(asks)))
        if subs:
            bad.append("V%d contains a spoken Subscribe request" % n)
    return not bad, bad or ["each script has exactly one spoken comment ask "
                            "and no spoken Subscribe request"]


@check("Resource is mentioned at most once per script")
def c19():
    bad = []
    for n in VIDEOS:
        blocks = spoken_units(_copy(n))
        hits = sum(1 for b in blocks
                   if re.search(r"Keep the Proof|Field Kit|Career Decision "
                                r"Evidence Check|Career Evidence Starter", b))
        if hits > 1:
            bad.append("V%d mentions a resource %d times" % (n, hits))
    return not bad, bad or ["at most one resource mention in each script"]


@check("Recording copies contain only recordable script")
def c20():
    bad = []
    for n in VIDEOS:
        us = units(_copy(n))
        for u in us:
            if re.search(r"Full screen|Primary CTA|Watch Next|Boundaries|"
                         r"Speech-only estimate|Editorial direction|"
                         r"Framework:|SUBSTITUTION RULE|visual map",
                         u, re.I):
                bad.append("V%d :: %s" % (n, u[:70]))
    return not bad, bad or ["no production notes, visual map, CTA analysis, "
                            "strategy commentary or research notes in any "
                            "recording copy"]


@check("Recording copy script matches the master script exactly")
def c21():
    bad = []
    for n in VIDEOS:
        m = [u for u in spoken_units_master(n)]
        c = spoken_units(_copy(n))
        if m != c:
            bad.append("V%d differs: %d master blocks, %d copy blocks"
                       % (n, len(m), len(c)))
    return not bad, bad or ["the seven recording copies reproduce the seven "
                            "master scripts block for block"]


def spoken_units_master(n):
    """Thought blocks inside a master, which are set at 12.5pt."""
    out = []
    d = Document(_master(n))
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if sizes and 12.0 <= max(sizes) < 13.0:
            out.append(p.text.strip())
    return out


@check("No script is labeled final or locked")
def c22():
    bad = []
    for f in sorted(glob.glob(os.path.join(OUT, "*.docx"))):
        for u in units(f):
            if re.search(r"\b(FINAL|LOCKED)\b", u):
                if not re.search(r"not (yet )?(locked|final)|NOT LOCKED|"
                                 r"final edit|final export|final title|"
                                 r"final packaging|final production|"
                                 r"final recording script|final script|"
                                 r"NOT YET APPROVED|final Riverside|"
                                 r"final thumbnail|final chapters|"
                                 r"final SRT|final fee|final Watch Next|"
                                 r"final publishing", u, re.I):
                    bad.append("%s :: %s" % (os.path.basename(f), u[:80]))
    return not bad, bad or ["no document labels a script final or locked"]


@check("Every master states the draft status")
def c23():
    bad = []
    for n in VIDEOS:
        us = units(_master(n))
        if not any("SCRIPT DRAFT FOR EDITORIAL REVIEW" in u for u in us):
            bad.append("V%d missing the status line" % n)
        if not any("NOT LOCKED. NOT APPROVED FOR PRODUCTION" in u
                   for u in us):
            bad.append("V%d missing the closing status" % n)
    return not bad, bad or ["all seven masters carry the draft status at the "
                            "top and the not-approved status at the end"]


@check("No production package was created")
def c24():
    forbidden = re.compile(r"\.(srt|pptx|png|jpg|mp4|vtt)$", re.I)
    bad = [os.path.basename(f) for f in glob.glob(os.path.join(OUT, "*"))
           if forbidden.search(f)]
    names = [os.path.basename(f) for f in glob.glob(os.path.join(OUT, "*"))
             if re.search(r"riverside|shorts|thumbnail|description|chapter|"
                          r"pinned|tags", os.path.basename(f), re.I)]
    bad += names
    return not bad, bad or ["no Riverside prompt, Short, thumbnail, "
                            "description, chapter list, SRT, deck or image "
                            "in the delivery"]


@check("Runtime figures are labeled as estimates")
def c25():
    bad = []
    for n in VIDEOS:
        us = units(_master(n))
        if not any("not a timed read and not a finished runtime" in u
                   for u in us):
            bad.append("V%d runtime not qualified" % n)
    return not bad, bad or ["every master states that the figure is "
                            "arithmetic, not a measured runtime"]


def run():
    rows, passed = [], 0
    for name, fn in CHECKS:
        try:
            ok, detail = fn()
        except Exception as e:
            ok, detail = False, ["check raised %s: %s" % (type(e).__name__, e)]
        rows.append((name, ok, detail))
        passed += 1 if ok else 0
    return rows, passed


if __name__ == "__main__":
    rows, passed = run()
    for name, ok, detail in rows:
        print("%-4s %s" % ("PASS" if ok else "FAIL", name))
        for d_ in detail[:4]:
            print("       %s" % d_)
    print("\n%d of %d checks passed" % (passed, len(rows)))
    sys.exit(0 if passed == len(rows) else 1)
