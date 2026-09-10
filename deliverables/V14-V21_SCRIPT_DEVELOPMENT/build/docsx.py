# -*- coding: utf-8 -*-
"""Document helpers for the V14 to V21 script-development deliverables.

Reuses the helpers built for the editorial-development batch, which are in
turn built on the shared house helpers. Nothing shared is edited here.
"""
import os, sys
DELIV = "/home/user/temidayoafonja-site/deliverables/"
sys.path.insert(0, DELIV + "V14-V21_EDITORIAL_DEVELOPMENT/build")

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docs1421 import (base_doc, para, title_block, rule, h, kv, callout,
                      table, bullets, sub, field, page_break, footer_note,
                      NAVY, GOLD, DIM, RED, _insert_ordered)

TW = 6.7


def block(d, text, size=12.5):
    """One recordable thought block.

    Set larger than body copy and given real space around it, because these
    are read from a screen at arm's length and delivered in one take.
    """
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.34
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def section_label(d, text):
    """A working label for the recording desk. Never spoken."""
    return para(d, text.upper(), size=8.5, bold=True, color=GOLD,
                before=16, after=8, keep=True)


def notspoken(d, text):
    return para(d, text, size=9, italic=True, color=DIM, after=10)


def numbered(d, items, size=10.5):
    for i, it in enumerate(items, 1):
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.34)
        p.paragraph_format.first_line_indent = Inches(-0.34)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.16
        r = p.add_run("%d.   " % i)
        r.bold = True
        r.font.size = Pt(size)
        r2 = p.add_run(it)
        r2.font.size = Pt(size)
