# -*- coding: utf-8 -*-
"""Document helpers for the Videos 14 to 21 final production packages."""
import os, sys
DELIV = "/home/user/temidayoafonja-site/deliverables/"
sys.path.append(DELIV + "V14-V21_EDITORIAL_DEVELOPMENT/build")

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docs1421 import (base_doc, para, title_block, rule, h, kv, callout,
                      table, bullets, sub, field, page_break, footer_note,
                      NAVY, GOLD, DIM, RED)

TW = 6.7


def block(d, text, size=12.5, after=14):
    """One recordable thought block, with its internal line breaks kept.

    The masters use line breaks inside a block to shape delivery. Those are
    performance instructions, so they are preserved rather than reflowed.
    """
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.34
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i:
            p.runs[-1].add_break()
        r = p.add_run(line)
        r.font.size = Pt(size)
    return p


def caption(d, text, size=9.5):
    """A quiet note under a block.

    sub() is a short gold label set in capitals. A full sentence passed to
    it renders as a paragraph of shouting, so anything sentence-shaped
    belongs here instead.
    """
    return para(d, text, size=size, color=DIM, before=5, after=10)


def section_label(d, text):
    """A working label from the master. Never spoken."""
    return para(d, text.upper(), size=8.5, bold=True, color=GOLD,
                before=16, after=8, keep=True)


def notspoken(d, text):
    return para(d, text, size=9, italic=True, color=DIM, after=10)


def numbered(d, items, size=10.5):
    for i, it in enumerate(items, 1):
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.34)
        p.paragraph_format.first_line_indent = Inches(-0.34)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.16
        r = p.add_run("%d.   " % i)
        r.bold = True
        r.font.size = Pt(size)
        r2 = p.add_run(it)
        r2.font.size = Pt(size)


def mono(path, lines):
    with open(path, "w") as f:
        f.write("\n".join(lines).rstrip() + "\n")
    return path


def hr(width=78):
    return "-" * width


def head(title, width=78):
    return ["=" * width, title, "=" * width, ""]
