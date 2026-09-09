# -*- coding: utf-8 -*-
"""The six September 9 locked final Recording Masters for Videos 8 to 13.

These files are the primary source of truth. This module only reads them, so
no spoken line can drift: every reading copy, trigger map and word count in the
batch is generated from the locked text rather than transcribed.

The spoken section begins after RECORDING SCRIPT STARTS HERE and ends before
END OF SPOKEN SCRIPT. Section headings, production labels and the closing
instruction about not reading the line are not spoken dialogue.
"""
import os, re, hashlib
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.abspath(os.path.join(HERE, "..", "_source"))
VIDEOS = (8, 9, 10, 11, 12, 13)

FILE = {n: "Video_%d_FINAL_Revised_Recording_Master.docx" % n for n in VIDEOS}
START = "RECORDING SCRIPT STARTS HERE"
END = "END OF SPOKEN SCRIPT"

# A production label rather than speech. These carry a "|" or are all caps.
LABEL = re.compile(r"^[A-Z0-9][A-Z0-9 ,:?+/'&.\-]*(\|.*)?$")


def path(n):
    return os.path.join(SRC, FILE[n])


def sha256(n):
    return hashlib.sha256(open(path(n), "rb").read()).hexdigest()


def _rows(n):
    return [p.text.strip() for p in Document(path(n)).paragraphs if p.text.strip()]


def header(n):
    """The master's own cover table, as a dict."""
    out = {}
    for t in Document(path(n)).tables:
        for r in t.rows:
            c = [x.text.strip() for x in r.cells]
            if len(c) == 2 and c[0]:
                out[c[0].upper()] = c[1]
    return out


def title(n):
    return _rows(n)[3]


def sections(n):
    """[(production label, [spoken lines]), ...] from the locked script only."""
    rows = _rows(n)
    i = rows.index(START) + 1
    j = rows.index(END)
    out, cur = [], None
    for t in rows[i:j]:
        is_label = ("|" in t and t.split("|")[0].strip().isupper()) or (
            t.isupper() and len(t) < 70) or re.match(r"^\d+ \| ", t)
        if is_label:
            cur = (t, [])
            out.append(cur)
        else:
            if cur is None:
                cur = ("", [])
                out.append(cur)
            cur[1].append(t)
    return [(a, b) for a, b in out if b]


def speech(n):
    return [t for _, body in sections(n) for t in body]


def words(n):
    return sum(len(t.split()) for t in speech(n))


def contains(n, sentence):
    """Is this exact sentence in the locked spoken script?"""
    def norm(s):
        return (s.replace("’", "'").replace("‘", "'")
                 .replace("“", '"').replace("”", '"'))
    return norm(sentence).strip() in norm("\n".join(speech(n)))


def estimate(n):
    """Speech-only estimate at the established 130 to 145 wpm band."""
    w = words(n)

    def mmss(m):
        return "%d:%02d" % (int(m), round((m - int(m)) * 60))
    return w, mmss(w / 145.0), mmss(w / 130.0)


if __name__ == "__main__":
    for n in VIDEOS:
        w, fast, slow = estimate(n)
        h = header(n)
        print("V%-3d %s" % (n, title(n)))
        print("     file    %s" % FILE[n])
        print("     sha256  %s" % sha256(n))
        print("     spoken  %d paragraphs, %d words, %s to %s at 130-145 wpm"
              % (len(speech(n)), w, fast, slow))
        print("     printed target: %s" % h.get("TARGET LENGTH"))
        print("     thumbnail: %s | resource: %s" % (h.get("THUMBNAIL"), h.get("RESOURCE")))
        print("     watch next: %s" % h.get("WATCH NEXT"))
        print("     sections: %d" % len(sections(n)))
        for lab, body in sections(n):
            print("        %-52s %d" % (lab[:52], len(body)))
        print()
