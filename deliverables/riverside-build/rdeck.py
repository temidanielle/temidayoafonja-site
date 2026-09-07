"""
Riverside edit-asset renderer.

Video visuals, not presentation slides. A Riverside asset is a full-frame
1920 x 1080 card that appears for a few seconds over a talking-head edit and
then gets out of the way, so the type is large, the copy is short, and the
important content sits inside caption-safe margins.

Two backends share one absolutely-positioned element list:

  render_pptx  -> editable PowerPoint, 13.333in x 7.5in
  render_html  -> HTML that Chromium rasterises to exact 1920 x 1080 PNGs

Brand values are the approved Capability Formation system: deep navy #112345
(the value the standing brand rules and the roadmap specify), the approved
cream and muted gold, Montserrat for display and DM Sans for body.
"""
import os
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
FONT_DIR = os.path.join(REPO, "fonts")
TTF_DIR = "/root/.fonts"

# ---------------------------------------------------------------- geometry
W, H = 1920, 1080
PX = 6350                        # EMU per design pixel (12192000 / 1920)
def E(px): return Emu(int(round(px * PX)))
def P(px): return Pt(px / 2.0)   # 1920px == 13.333in == 144 px/inch

MARGIN = 160                     # side margin, generous for video
CW = W - 2 * MARGIN              # 1600 content width
TOP = 150                        # first baseline zone
CAPTION_SAFE = 880               # nothing important below this line

# ------------------------------------------------------------------ colour
NAVY      = RGBColor(0x11, 0x23, 0x45)
CREAM     = RGBColor(0xF5, 0xF1, 0xE8)
GOLD      = RGBColor(0xC9, 0xA8, 0x4C)
RUST      = RGBColor(0xC1, 0x44, 0x0E)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CREAM_DIM = RGBColor(0xB9, 0xC3, 0xD2)
NAVY_DIM  = RGBColor(0x5A, 0x6B, 0x82)
RULE_CREAM = RGBColor(0xE2, 0xDC, 0xCC)
RULE_NAVY  = RGBColor(0x2A, 0x3C, 0x5C)

DISPLAY = "Montserrat"
BODY    = "DM Sans"

def hexc(c): return "#%s" % str(c)

# TTF files used only for width measurement in QA
TTF = {
    (DISPLAY, True):  os.path.join(TTF_DIR, "Montserrat-Bold.ttf"),
    (DISPLAY, False): os.path.join(TTF_DIR, "Montserrat-Regular.ttf"),
    (BODY, True):     os.path.join(TTF_DIR, "DMSans-Bold.ttf"),
    (BODY, False):    os.path.join(TTF_DIR, "DMSans-Regular.ttf"),
}


class Card(object):
    """One Riverside asset: filename, purpose metadata and drawing elements."""

    def __init__(self, index, filename):
        self.index, self.filename = index, filename
        self.els = []
        self.notes = ""

    def add(self, el):
        self.els.append(el)
        return el


DEFAULT_PARA = dict(size=40, font=DISPLAY, color=NAVY, bold=False, italic=False,
                    align="l", spacing=1.14, tracking=0, space_before=0,
                    space_after=0)


# --------------------------------------------------------------- primitives
def rect(c, x, y, w, h, fill=None, line=None, lw=2, shape="rect"):
    return c.add(dict(t="rect", x=x, y=y, w=w, h=h, fill=fill, line=line,
                      lw=lw, shape=shape, dash=None))


def block(c, x, y, w, lines, anchor="t", h=None):
    paras = []
    for txt, st in lines:
        p = dict(DEFAULT_PARA)
        p.update(st)
        p["text"] = txt
        paras.append(p)
    return c.add(dict(t="text", x=x, y=y, w=w, h=h or 0, anchor=anchor,
                      paras=paras))


def bg(c, color):
    return rect(c, 0, 0, W, H, fill=color)


def eyebrow(c, x, y, text, color=GOLD, size=30, w=None, align="l"):
    return block(c, x, y, w or CW, [(text.upper(),
                 dict(size=size, font=DISPLAY, color=color, bold=True,
                      tracking=5.0, align=align, spacing=1.0))])


def rule(c, x, y, w, color=GOLD, h=4):
    return rect(c, x, y, w, h, fill=color)


def logomark(c, x, y, unit=34, gap=9, gold=GOLD, rust=RUST, lw=3):
    """The Capability Formation mark: four squares, top-right filled."""
    for (cx, cy), filled in (((0, 0), False), ((1, 0), True),
                             ((0, 1), False), ((1, 1), False)):
        rect(c, x + cx * (unit + gap), y + cy * (unit + gap), unit, unit,
             fill=rust if filled else None,
             line=None if filled else gold, lw=lw)


# ------------------------------------------------------------ pptx backend
_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def _pptx_rect(sl, el):
    shape = MSO_SHAPE.OVAL if el["shape"] == "oval" else MSO_SHAPE.RECTANGLE
    s = sl.shapes.add_shape(shape, E(el["x"]), E(el["y"]), E(el["w"]), E(el["h"]))
    s.shadow.inherit = False
    if el["fill"] is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = el["fill"]
    if el["line"] is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = el["line"]
        s.line.width = E(el["lw"])
    return s


def _pptx_text(sl, el):
    h = el["h"] or 240
    tb = sl.shapes.add_textbox(E(el["x"]), E(el["y"]), E(el["w"]), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR[el["anchor"]]
    for i, sp in enumerate(el["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[sp["align"]]
        p.line_spacing = sp["spacing"]
        if sp["space_before"]:
            p.space_before = P(sp["space_before"])
        if sp["space_after"]:
            p.space_after = P(sp["space_after"])
        parts = sp["text"].split("\n")
        for j, line in enumerate(parts):
            r = p.add_run()
            r.text = line
            f = r.font
            f.name, f.size = sp["font"], P(sp["size"])
            f.bold, f.italic = sp["bold"], sp["italic"]
            f.color.rgb = sp["color"]
            rPr = r.font._rPr
            if sp["tracking"]:
                rPr.set("spc", str(int(sp["tracking"] * 50)))
            if sp.get("strike"):
                rPr.set("strike", "sngStrike")
            for tag in ("a:ea", "a:cs"):
                rPr.append(rPr.makeelement(qn(tag), {"typeface": sp["font"]}))
            if j < len(parts) - 1:
                p._p.append(p._p.makeelement(qn("a:br"), {}))
    return tb


def render_pptx(cards, path):
    from pptx import Presentation
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    for cd in cards:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        for el in cd.els:
            {"rect": _pptx_rect, "text": _pptx_text}[el["t"]](sl, el)
        if cd.notes:
            sl.notes_slide.notes_text_frame.text = cd.notes
    prs.save(path)
    return path


# ------------------------------------------------------------ html backend
_FONT_FACES = [
    ("Montserrat", 400, "normal", "Montserrat-400-normal-latin-49e242.woff2"),
    ("Montserrat", 500, "normal", "Montserrat-500-normal-latin-49e242.woff2"),
    ("Montserrat", 600, "normal", "Montserrat-600-normal-latin-49e242.woff2"),
    ("Montserrat", 700, "normal", "Montserrat-700-normal-latin-49e242.woff2"),
    ("DM Sans", 400, "normal", "DMSans-400-normal-latin-1c49a6.woff2"),
    ("DM Sans", 500, "normal", "DMSans-500-normal-latin-1c49a6.woff2"),
    ("DM Sans", 700, "normal", "DMSans-600-normal-latin-1c49a6.woff2"),
]

_CSS = """
* { margin:0; padding:0; box-sizing:border-box; }
html,body { background:#ffffff; }
.slide { position:relative; width:1920px; height:1080px; overflow:hidden;
         background:#F5F1E8; }
.el { position:absolute; }
.tx { display:flex; flex-direction:column; }
.tx.t { justify-content:flex-start; }
.tx.m { justify-content:center; }
.tx.b { justify-content:flex-end; }
p { white-space:pre-wrap; }
"""


def _esc(s):
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
             .replace("\n", "<br>"))


def _html_el(el, rel):
    x, y, w, h = el["x"], el["y"], el["w"], el["h"]
    if el["t"] == "rect":
        st = ["left:%gpx" % x, "top:%gpx" % y, "width:%gpx" % w, "height:%gpx" % h]
        if el["fill"] is not None:
            st.append("background:%s" % hexc(el["fill"]))
        if el["line"] is not None:
            st.append("border:%gpx solid %s" % (el["lw"], hexc(el["line"])))
        if el["shape"] == "oval":
            st.append("border-radius:50%")
        return '<div class="el" style="%s"></div>' % ";".join(st)

    st = ["left:%gpx" % x, "top:%gpx" % y, "width:%gpx" % w]
    if h:
        st.append("height:%gpx" % h)
    out = ['<div class="el tx %s" style="%s">' % (el["anchor"], ";".join(st))]
    for sp in el["paras"]:
        ps = ["font-family:'%s',sans-serif" % sp["font"],
              "font-size:%gpx" % sp["size"],
              "font-weight:%d" % (700 if sp["bold"] else 400),
              "line-height:%g" % sp["spacing"],
              "color:%s" % hexc(sp["color"]),
              "text-align:%s" % {"l": "left", "c": "center", "r": "right"}[sp["align"]]]
        if sp["italic"]:
            ps.append("font-style:italic")
        if sp.get("strike"):
            ps.append("text-decoration:line-through")
            ps.append("text-decoration-thickness:3px")
        if sp["tracking"]:
            ps.append("letter-spacing:%gpx" % sp["tracking"])
        if sp["space_before"]:
            ps.append("margin-top:%gpx" % sp["space_before"])
        if sp["space_after"]:
            ps.append("margin-bottom:%gpx" % sp["space_after"])
        out.append('<p style="%s">%s</p>' % (";".join(ps), _esc(sp["text"])))
    out.append("</div>")
    return "".join(out)


def render_html(cards, path, title="Riverside assets"):
    outdir = os.path.dirname(os.path.abspath(path))
    rel = lambda p: os.path.relpath(p, outdir).replace(os.sep, "/")
    faces = "\n".join(
        "@font-face{font-family:'%s';font-weight:%d;font-style:%s;"
        "src:url('%s') format('woff2');font-display:block;}"
        % (fam, wt, sty, rel(os.path.join(FONT_DIR, fn)))
        for fam, wt, sty, fn in _FONT_FACES)
    body = "\n".join(
        '<div class="slide" id="s%d">%s</div>'
        % (i + 1, "".join(_html_el(el, rel) for el in cd.els))
        for i, cd in enumerate(cards))
    html = ("<!doctype html><html><head><meta charset='utf-8'><title>%s</title>"
            "<style>%s\n%s</style></head><body>%s</body></html>"
            % (title, faces, _CSS, body))
    with open(path, "w") as f:
        f.write(html)
    return path


def shoot(html_path, png_dir, names):
    """Rasterise each card to an exact 1920 x 1080 PNG."""
    from playwright.sync_api import sync_playwright
    CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
    os.makedirs(png_dir, exist_ok=True)
    made = []
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path=CHROME, args=["--no-sandbox"])
        pg = b.new_page(viewport={"width": W, "height": H}, device_scale_factor=1)
        pg.goto("file://" + os.path.abspath(html_path), wait_until="load")
        pg.wait_for_timeout(800)
        for i, name in enumerate(names, start=1):
            p = os.path.join(png_dir, name)
            pg.query_selector("#s%d" % i).screenshot(path=p)
            made.append(p)
        b.close()
    return made


# ------------------------------------------------------- text measurement
_FONT_CACHE = {}


def _pil_font(name, bold, size):
    key = (name, bool(bold), int(size))
    if key not in _FONT_CACHE:
        from PIL import ImageFont
        _FONT_CACHE[key] = ImageFont.truetype(TTF[(name, bool(bold))], int(size))
    return _FONT_CACHE[key]


def wrap_lines(text, width, font=DISPLAY, size=40, bold=False, tracking=0):
    """Greedy wrap at `width` px, honouring explicit newlines. Returns lines."""
    f = _pil_font(font, bold, size)

    def measure(s):
        return f.getlength(s) + tracking * max(len(s) - 1, 0)

    out = []
    for hard in text.split("\n"):
        words, cur = hard.split(" "), ""
        for w in words:
            trial = w if not cur else cur + " " + w
            if measure(trial) <= width or not cur:
                cur = trial
            else:
                out.append(cur)
                cur = w
        out.append(cur)
    return out


def text_height(text, width, font=DISPLAY, size=40, bold=False, spacing=1.14,
                tracking=0):
    n = len(wrap_lines(text, width, font, size, bold, tracking))
    return n * size * spacing


def max_line_width(text, width, font=DISPLAY, size=40, bold=False, tracking=0):
    f = _pil_font(font, bold, size)
    lines = wrap_lines(text, width, font, size, bold, tracking)
    return max(f.getlength(l) + tracking * max(len(l) - 1, 0) for l in lines)
