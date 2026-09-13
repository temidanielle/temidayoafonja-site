# -*- coding: utf-8 -*-
"""Document helpers for the V22 and V23 production packages.

The house helpers are imported, never edited, so every earlier batch's
documents still build byte for byte. What is added here is local to these two
packages.
"""
import os, sys
DELIV = "/home/user/temidayoafonja-site/deliverables/"
sys.path.append(DELIV + "VIDEOS_4-21_CORRECTED_RUNTIME/build")

from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docs421f import (base_doc, para, title_block, rule, h, kv, callout,
                      table, bullets, sub, field, page_break, footer_note,
                      block, caption, section_label, notspoken, numbered,
                      mono, hr, head, NAVY, GOLD, DIM, RED, TW)


def spoken(d, text):
    """One recordable paragraph of the approved script, unchanged."""
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(11)
    p.paragraph_format.line_spacing = 1.34
    r = p.add_run(text)
    r.font.size = Pt(12.5)
    return p


def marker(d, code, label):
    """An editorial navigation marker from the script. Never a runtime."""
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    pr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), "6")
    bt.set(qn("w:color"), "8A6D1E")
    b.append(bt)
    pr.append(b)
    for tag in ("keepNext", "keepLines"):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "1")
        pr.append(e)
    r = p.add_run(code + "   ")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = DIM
    r2 = p.add_run(label.upper())
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = NAVY
    return p


def cue(d, mode, key, detail):
    """A production cue inside a recording document. Never spoken."""
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.left_indent = Inches(0.15)
    pr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    lt = OxmlElement("w:left")
    lt.set(qn("w:val"), "single")
    lt.set(qn("w:sz"), "12")
    lt.set(qn("w:color"), "C9A84C")
    lt.set(qn("w:space"), "8")
    b.append(lt)
    pr.append(b)
    r = p.add_run("[ %s ]  " % mode.upper())
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = GOLD
    r2 = p.add_run(key)
    r2.bold = True
    r2.font.size = Pt(9)
    r2.font.color.rgb = NAVY
    if detail:
        r3 = p.add_run("   " + detail)
        r3.font.size = Pt(9)
        r3.font.color.rgb = DIM
    return p
