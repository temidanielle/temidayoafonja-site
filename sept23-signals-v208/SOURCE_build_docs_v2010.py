# -*- coding: utf-8 -*-
"""SOP v2.0.11 -> v2.0.12 and QA report v2.0.12 -> v2.0.13, for deck v2.0.10.

Nothing in the timeline changed, so the run of show keeps every time. What
changed is which slide the room sees at 50:00, the slide count, and the record
of the six decisions applied in v2.0.10.

The QA report also records the one deliberate exception to the banned-word rule:
"genuinely" survives inside a locked scored statement, and removing it would
change the instrument.
"""
import copy, hashlib, os, re, shutil
import docx
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
SOP_SRC = os.path.join(
    HERE,
    "Capability_Formation_Free_Flagship_SOP_Sept23_2026_60MIN_v2.0.11_CANDIDATE.docx")
SOP_DST = os.path.join(
    HERE,
    "Capability_Formation_Free_Flagship_SOP_Sept23_2026_60MIN_v2.0.12_CANDIDATE.docx")
RPT_SRC = os.path.join(HERE,
                       "Free_Flagship_60MIN_v2.0.12_Change_Log_and_QA_Report.docx")
RPT_DST = os.path.join(HERE,
                       "Free_Flagship_60MIN_v2.0.13_Change_Log_and_QA_Report.docx")

EM = "—"
OLD_STAMP = "Friday, September 25, 2026 at 9:10 AM CT"
NEW_STAMP = "Friday, September 25, 2026 at 11:40 AM CT"

SOP_FIXES = [
 ("advance to the Q&A slide.", "advance to slide 29, Questions & Applications.", 1),
 ("Q&A slide", "29", 1),
 ("Confirm aloud that recording has stopped. Method questions only. Hard stop "
  "at 60:00.",
  "Advance to slide 29, which says LIVE ONLY and RECORDING OFF on its face. "
  "Confirm aloud that recording has stopped. Method questions only. Hard stop "
  "at 60:00.", 1),
]

SECTION_HEAD = "v2.0.10 decisions"
SECTION_BODY = (
 "Six owner decisions applied to deck v2.0.9. The timeline did not change, so "
 "the run of show keeps every time in it. The workbook stays at v2.0.3 because "
 "nothing in these decisions touches it.")

CHANGES = [
 ("Where", "Change"),
 ("Timing",
  "Seventy-five seconds each for the two definition slides is accepted. No "
  "other timing changed, so this build changes no time at all."),
 ("Three constructions in notes",
  "The not-X-it-is-Y constructions on the Signals and Keep Working Privately "
  "notes are replaced with the owner's wording, verbatim. A participant who "
  "sees several signals at once is now described by where they are, not by "
  "where they are not. The Field Kit is described by what it is. The retired "
  "offer is handled with an instruction, Leave it unmentioned, rather than with "
  "three denials."),
 ("Banned words",
  "actually and genuinely removed from nine note instances and from one "
  "appendix face, which now reads What translation means. Sentences were "
  "rewritten where the removal left them awkward. The owner's brief said twelve "
  "note instances; v2.0.9 carried nine, because three of the twelve sat in the "
  "Definitions and Matrix divider notes that v2.0.9 had already replaced."),
 ("ONE DELIBERATE EXCEPTION",
  "genuinely is KEPT in the scored statement \"Looking back six months, the work "
  "I do now would have been genuinely hard for me then.\" That statement is part "
  "of the locked twelve-statement instrument and is carried verbatim by the "
  "workbook. Removing the word would change the instrument, not the prose, and "
  "the instrument is not open. Any future banned-word sweep must skip it."),
 ("Appositives",
  "The short X, not Y appositives are approved house voice and were not "
  "touched: Direction, not a plan. Questions to investigate, not predictions. "
  "Phrases, not essays. Privately, and provisionally."),
 ("Q&A",
  "The real Questions & Applications slide, the one whose face says LIVE ONLY "
  "and RECORDING OFF, is unhidden and is now slide 29. The bare Q & A divider "
  "is hidden. This matters operationally: the recording-off instruction is on "
  "the face of the slide the room sees at 50:00, not on a divider."),
 ("Poll slides",
  "The appendix version is kept, with its appendix eyebrow and its anonymity "
  "line. The other is deleted. The deck goes from 43 slides to 42."),
 ("Footer numbers",
  "Renumbered again, because unhiding Questions & Applications added a numbered "
  "slide to the visible core. Twenty-nine numbered core slides."),
]

NEW_CHECKS = [
 ("259", "Banned words", "PASS",
  "One instance of actually, genuinely or honestly remains in the whole deck, "
  "faces and notes, and it is the locked scored statement. Asserted by the "
  "build, which fails if the survivor is anywhere else."),
 ("260", "Banned-word exception", "PASS",
  "The survivor is recorded as a deliberate exception tied to the instrument, "
  "not an oversight."),
 ("261", "Constructions", "PASS",
  "None of the three not-X-it-is-Y strings survives, and all three replacements "
  "are present verbatim."),
 ("262", "Appositives", "PASS",
  "The approved X, not Y appositives are unchanged."),
 ("263", "Q&A visibility", "PASS",
  "Questions & Applications is visible and carries LIVE ONLY and RECORDING OFF "
  "on its face. Every bare Q & A slide is hidden."),
 ("264", "Poll slides", "PASS",
  "Exactly one remains, hidden, and it is the one with the appendix eyebrow and "
  "the anonymity line."),
 ("265", "No regression", "PASS",
  "Zero em dashes, zero British spellings, both retired category names absent, "
  "and the opening block, the two seventy-five second definitions and the "
  "protected twelve minutes all read as v2.0.9 left them."),
 ("266", "Slide count", "PASS", "43 to 42, one deletion, nothing else removed."),
]


def carrier(para):
    runs = para.runs
    return next((r for r in runs if r.text.strip()), runs[0]) if runs else None


def set_para(para, text):
    keep = carrier(para)
    assert keep is not None
    keep.text = text
    for r in para.runs:
        if r._r is not keep._r:
            r._r.getparent().remove(r._r)


def set_cell(cell, text):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    set_para(p, text)


def every_run(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            yield r
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        yield r


def body_text(doc):
    return ("\n".join(p.text for p in doc.paragraphs) + "\n"
            + "\n".join(c.text for t in doc.tables for row in t.rows
                        for c in row.cells))


def count_em(doc):
    return body_text(doc).count(EM)


def find_para(doc, needle):
    hits = [i for i, p in enumerate(doc.paragraphs) if needle in p.text]
    assert len(hits) == 1, f"{needle!r} found {len(hits)} times"
    return hits[0]


def swap(doc, pairs):
    hits = {old: 0 for old, _, _ in pairs}
    for r in every_run(doc):
        t = r.text
        for old, new, _ in pairs:
            if old in t:
                hits[old] += t.count(old)
                t = t.replace(old, new)
        if t != r.text:
            r.text = t
    for old, _, want in pairs:
        assert hits[old] == want, f"{old[:44]!r}: {hits[old]} not {want}"


def build_sop():
    shutil.copyfile(SOP_SRC, SOP_DST)
    d = docx.Document(SOP_DST)
    swap(d, SOP_FIXES)
    for r in every_run(d):
        t = r.text.replace("v2.0.11", "v2.0.12").replace(OLD_STAMP, NEW_STAMP)
        if t != r.text:
            r.text = t
    d.save(SOP_DST)

    out = docx.Document(SOP_DST)
    b = body_text(out)
    assert count_em(out) == 0
    assert "v2.0.12" in b and "v2.0.11" not in b
    assert "LIVE ONLY and RECORDING OFF" in b
    assert "18:00–30:00" in b, "the protected block moved"
    print("built", os.path.basename(SOP_DST))
    print("  sha256", hashlib.sha256(open(SOP_DST, "rb").read()).hexdigest())


def build_report():
    shutil.copyfile(RPT_SRC, RPT_DST)
    d = docx.Document(RPT_DST)

    # counts and stamps first, so the rows below can name v2.0.12 safely
    for r in every_run(d):
        t = (r.text.replace("258 of 258 PASS", "266 of 266 PASS")
                   .replace("258 items", "266 items")
                   .replace("v2.0.12", "v2.0.13")
                   .replace("PowerPoint v2.0.9", "PowerPoint v2.0.10")
                   .replace("SOP v2.0.11", "SOP v2.0.12")
                   .replace(OLD_STAMP, NEW_STAMP))
        if t != r.text:
            r.text = t

    # the v2.0.9 section is now history, so its opening line moves to past tense
    swap(d, [("The deck is v2.0.9, the workbook is v2.0.3 and the SOP is v2.0.11.",
              "At that revision the deck was v2.0.9, the workbook v2.0.3 and the "
              "SOP v2.0.11.", 1)])
    anchor = d.paragraphs[find_para(d, "Verification. ")]
    head_src = d.paragraphs[find_para(d, "Signals Worth Watching and the v2.0.9")]
    body_src = d.paragraphs[find_para(d, "A new section and six owner decisions")]

    head_el = copy.deepcopy(head_src._element)
    anchor._element.addprevious(head_el)
    set_para(Paragraph(head_el, anchor._parent), SECTION_HEAD)
    body_el = copy.deepcopy(body_src._element)
    head_el.addnext(body_el)
    set_para(Paragraph(body_el, anchor._parent), SECTION_BODY)

    tbl_el = copy.deepcopy(d.tables[9]._element)
    body_el.addnext(tbl_el)
    new_tbl = [t for t in d.tables if t._element is tbl_el][0]
    template = copy.deepcopy(new_tbl.rows[1]._element)
    while len(new_tbl.rows) > 1:
        new_tbl._element.remove(new_tbl.rows[-1]._element)
    for ci, v in enumerate(CHANGES[0]):
        set_cell(new_tbl.rows[0].cells[ci], v)
    for row in CHANGES[1:]:
        new_tbl._element.append(copy.deepcopy(template))
        for ci, v in enumerate(row):
            set_cell(new_tbl.rows[-1].cells[ci], v)

    qa = d.tables[-1]
    assert [c.text for c in qa.rows[0].cells] == ["#", "Category", "Status",
                                                  "Notes"]
    qtemplate = copy.deepcopy(qa.rows[-1]._element)
    for row in NEW_CHECKS:
        qa._element.append(copy.deepcopy(qtemplate))
        for ci, v in enumerate(row):
            set_cell(qa.rows[-1].cells[ci], v)

    for r in every_run(d):
        if EM in r.text:
            r.text = re.sub(r"\.\s*\.", ".", re.sub(r"\s*" + EM + r"\s*", ". ",
                                                    r.text))
    d.save(RPT_DST)

    out = docx.Document(RPT_DST)
    b = body_text(out)
    assert count_em(out) == 0
    assert "266 of 266 PASS" in b and "258 of 258" not in b
    assert "266 items" in b
    assert SECTION_HEAD in b
    assert "PowerPoint v2.0.10, workbook v2.0.3, SOP v2.0.12" in b, \
        "the source-of-truth map is wrong"
    assert "ONE DELIBERATE EXCEPTION" in b, "the exception is not recorded"
    assert "genuinely hard for me then" in b, "the exception does not quote it"
    for n, _, _, _ in NEW_CHECKS:
        assert n in b
    print("built", os.path.basename(RPT_DST))
    print(f"  new change rows: {len(CHANGES) - 1}   new checks: {len(NEW_CHECKS)}")
    print("  sha256", hashlib.sha256(open(RPT_DST, "rb").read()).hexdigest())


if __name__ == "__main__":
    build_sop()
    build_report()
