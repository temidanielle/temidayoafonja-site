#!/usr/bin/env python3
"""Build 03_YOUR_PROFESSIONAL_RECORD.docx (primary copyable working document)
and 03_YOUR_PROFESSIONAL_RECORD.md (portable plain-text mirror)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
import record_model as M

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = "/home/user/temidayoafonja-site"
OUT_DOCX = f"{ROOT}/_build/keep-the-proof-v2/production/bundle/03_YOUR_PROFESSIONAL_RECORD.docx"
# The .md is an INTERNAL source/backup mirror only — it is NOT shipped in the
# customer bundle (customers get Word or the fillable PDF, not a third format).
OUT_MD = f"{ROOT}/_build/keep-the-proof-v2/production/source/03_YOUR_PROFESSIONAL_RECORD.md"

NAVY = RGBColor(0x11, 0x23, 0x45)
GOLD = RGBColor(0xA8, 0x86, 0x2E)   # slightly deeper gold for text legibility on white
GREY = RGBColor(0x6b, 0x70, 0x7d)
INK = RGBColor(0x1c, 0x23, 0x33)

doc = Document()
# base style
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for s in doc.sections:
    s.top_margin = Inches(0.9); s.bottom_margin = Inches(0.9)
    s.left_margin = Inches(1.0); s.right_margin = Inches(1.0)

# Footer page number (centered), applied to the whole document
def add_footer_page_number():
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    run.font.size = Pt(9); run.font.color.rgb = GREY
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); run._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt
            run._r.append(it)
add_footer_page_number()

def para(text="", size=10.5, bold=False, italic=False, color=INK, after=6, before=0, align=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.color.rgb = color
    return p

ICON_DIR = f"{ROOT}/_build/keep-the-proof-v2/production/assets/icons"

def add_icon(p, name, size=0.16):
    """Prepend a small inline icon (with alt text) to a heading paragraph."""
    run = p.add_run()
    shape = run.add_picture(f"{ICON_DIR}/{name}.png", width=Inches(size))
    alt = __import__("icons").LABELS.get(name, name)
    docPr = shape._inline.docPr
    docPr.set("descr", alt); docPr.set("title", alt)
    sp = p.add_run("  "); sp.font.size = Pt(11)
    return p

def h_section(num, title, icon=None):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    if icon: add_icon(p, icon)
    r = p.add_run(f"Section {num}.  "); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = GOLD
    r2 = p.add_run(title); r2.bold = True; r2.font.size = Pt(14); r2.font.color.rgb = NAVY
    _bottom_border(p, "C9A84C", 8)

def h_block(title, color=NAVY, size=11.5, icon=None):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    if icon: add_icon(p, icon, size=0.14)
    r = p.add_run(title); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = color
    return p

def _bottom_border(p, hexcolor="BBBBBB", sz=6):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), hexcolor)
    pbdr.append(bottom); pPr.append(pbdr)

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def cell_borders(cell, edges):
    # edges: {'left': (size, color), ...}; unspecified edges set to nil
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement('w:tcBorders')
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement('w:' + edge)
        if edge in edges:
            sz, col = edges[edge]
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
            el.set(qn('w:space'), '0'); el.set(qn('w:color'), col)
        else:
            el.set(qn('w:val'), 'nil')
        tb.append(el)
    tcPr.append(tb)

def field(label, hint, lines):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(label); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NAVY
    if hint:
        rh = p.add_run("  " + hint); rh.italic = True; rh.font.size = Pt(8.5); rh.font.color.rgb = GREY
    for i in range(lines):
        lp = doc.add_paragraph(); lp.paragraph_format.space_before = Pt(9); lp.paragraph_format.space_after = Pt(0)
        _bottom_border(lp, "C9C4B7", 4)

def checkbox_line(text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    r = p.add_run("☐  "); r.font.size = Pt(11); r.font.color.rgb = GOLD
    r2 = p.add_run(text); r2.font.size = Pt(10)

def bullet(text_bold, text):
    p = doc.add_paragraph(style=None); p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("•  "); r.font.color.rgb = GOLD; r.bold = True
    if text_bold:
        rb = p.add_run(text_bold + " "); rb.bold = True; rb.font.size = Pt(10); rb.font.color.rgb = NAVY
    rt = p.add_run(text); rt.font.size = Pt(10)

def divider():
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)

# ---------- Title ----------
tp = para(M.TITLE, size=26, bold=True, color=NAVY, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
tp.runs[0].font.name = "Cambria"
sp = para(M.SUBTITLE, size=13, italic=True, color=GOLD, after=10)
sp.runs[0].font.name = "Cambria"

para(M.INTRO_LEAD, after=8)
for b, t in M.INTRO_BULLETS:
    bullet(b, t)
para(M.INTRO_SECTIONS, before=4, after=6)
para(M.FORMATS_NOTE, italic=True, color=GREY, after=4)

# ---------- Opener: Before you rebuild anything ----------
h_block("Before you rebuild anything", color=NAVY, size=12, icon="reconstruct")
para(M.THREE_MOMENTS_LABEL, size=10, bold=True, color=NAVY, before=2, after=1)
para(M.THREE_MOMENTS_HINT, italic=True, color=GREY, size=8.5, after=3)
for i in range(3):
    lp = doc.add_paragraph(); lp.paragraph_format.space_before = Pt(10); _bottom_border(lp, "C9C4B7", 4)
# Words That Hold, as a calm cream panel with a gold left border (matching the handbook)
para("", after=2)
wtab = doc.add_table(rows=1, cols=1)
wcell = wtab.rows[0].cells[0]
shade_cell(wcell, "F5F1E8")
cell_borders(wcell, {"left": (24, "C9A84C"), "top": (4, "E0D7C4"), "bottom": (4, "E0D7C4"), "right": (4, "E0D7C4")})
# replace the cell's default empty paragraph
wcell.paragraphs[0]._p.getparent().remove(wcell.paragraphs[0]._p)
def _cellp(sb=0, sa=3, indent=0.10):
    p = wcell.add_paragraph()
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.left_indent = Inches(indent); p.paragraph_format.right_indent = Inches(0.08)
    return p
_lp = _cellp(6, 5); add_icon(_lp, "words", size=0.14)
_r = _lp.add_run(M.WORDS_TITLE.upper()); _r.bold = True; _r.font.size = Pt(9); _r.font.color.rgb = GOLD
_r.font.name = "Calibri"
_ld = _cellp(0, 6); _rl = _ld.add_run(M.WORDS_LEAD); _rl.italic = True; _rl.font.size = Pt(8.5); _rl.font.color.rgb = GREY
for w in M.WORDS_AFFIRMATIONS:
    _wp = _cellp(0, 5); _rw = _wp.add_run(w); _rw.italic = True; _rw.font.size = Pt(12); _rw.font.name = "Cambria"; _rw.font.color.rgb = NAVY
_ff = _cellp(4, 2); _rf = _ff.add_run(M.WORDS_FILLIN); _rf.bold = True; _rf.font.size = Pt(10); _rf.font.color.rgb = NAVY
_fl = _cellp(2, 6); _bottom_border(_fl, "C9A84C", 4)
# Read It Back workspace, directly after the Words That Hold box
h_block("Read It Back", color=NAVY, size=12, icon="words")
for _p in ["A sentence that shrinks my work",
           "What would have been different if I had not been there?",
           "The line I am learning to believe",
           "The entries that support it"]:
    field(_p, "", 2)
divider()

# ---------- Section 1 ----------
h_section(1, "Capture Log", icon="capture")
para("Catch work the moment it happens, in two minutes, before the details soften. Fill one block per event; expand only the strongest into Full Entries. Duplicate the block below for each new capture.", after=6)
for n in range(1, 4):
    h_block(f"Capture {n}", color=GOLD, size=10.5)
    for label, hint, lines in M.CAPTURE_FIELDS:
        field(label, hint, lines)
    divider()
para("To add more, copy any Capture block above and paste it here.", italic=True, color=GREY)

# ---------- Section 2 ----------
h_section(2, "Full Entries")
para("For work worth keeping in depth, expand a capture into a full entry. Fill only the fields that apply; a strong entry with six good fields beats a padded one with sixteen. Duplicate the whole entry template for each new entry.", after=6)
for n in range(1, 3):
    h_block(f"Full Entry {n}", color=GOLD, size=11)
    for cluster_title, fields in M.FULL_ENTRY_CLUSTERS:
        h_block(cluster_title, color=NAVY, size=10.5)
        for label, hint, lines in fields:
            field(label, hint, lines)
    para(M.RECON_MARKER, italic=True, color=GREY, before=4, after=4)
    divider()
para("To add more, copy a Full Entry template above and paste it here.", italic=True, color=GREY)

# ---------- Section 3 ----------
h_section(3, "Corroboration list")
para("Keep, in one place, who could confirm your work while they still remember it. This is corroboration readiness, not networking: it records that a witness exists and what they saw. It does not decide any employer's choice.", after=6)
for n in range(1, 5):
    h_block(f"Corroboration {n}", color=GOLD, size=10.5)
    for label, hint, lines in M.CORROB_FIELDS:
        field(label, hint, lines)
    divider()
para(M.CORROB_NOTE, italic=True, color=GREY)

# ---------- Section 4 ----------
h_section(4, "Translation and Proof Line workspace", icon="proofline")
para("Turn internal language into portable language, and build the portable sentence you can reuse.", after=6)
h_block("Translation worksheet", color=NAVY, size=10.5)
para("Work down the eight moves and apply the ones that fit, then the protection rule.", after=4)
for m in M.TRANSLATION_MOVES:
    checkbox_line(m)
para(M.TRANSLATION_PROTECTION, italic=True, color=GREY, before=2, after=8)
h_block("Proof Line builder", color=NAVY, size=10.5)
para("Combine, in whatever order reads well, up to five ingredients into one portable sentence.", after=4)
for label, hint, lines in M.PROOFLINE_INGREDIENTS:
    field(label, hint, lines)
para("Your Proof Line", size=10, bold=True, color=NAVY, before=6, after=1)
for i in range(2):
    lp = doc.add_paragraph(); lp.paragraph_format.space_before = Pt(9); _bottom_border(lp, "C9C4B7", 4)
para(M.PROOFLINE_RULE, italic=True, color=GREY, before=2, after=8)
h_block("Proof Lines kept", color=NAVY, size=10.5)
para("A running list of your finished Proof Lines, ready to pull for a review, a resume, or a conversation.", after=4)
for i in range(6):
    lp = doc.add_paragraph(); lp.paragraph_format.space_before = Pt(9); _bottom_border(lp, "C9C4B7", 4)

# ---------- Section 5: Match Your Proof to a Role ----------
h_section(5, M.MATCH_TITLE, icon="match")
para(M.MATCH_INTRO, after=6)
mtable = doc.add_table(rows=1, cols=3); mtable.style = "Table Grid"
mhdr = mtable.rows[0].cells
for i, col in enumerate(M.MATCH_COLUMNS):
    mhdr[i].text = ""
    rp = mhdr[i].paragraphs[0].add_run(col); rp.bold = True; rp.font.size = Pt(9); rp.font.color.rgb = NAVY
# example row (italic grey)
exrow = mtable.add_row().cells
for i, val in enumerate(M.MATCH_EXAMPLE):
    exrow[i].text = ""
    pfx = exrow[i].paragraphs[0].add_run("Example.  " if i == 0 else "")
    pfx.bold = True; pfx.font.size = Pt(8.5); pfx.font.color.rgb = GOLD
    rr = exrow[i].paragraphs[0].add_run(val); rr.italic = True; rr.font.size = Pt(9); rr.font.color.rgb = GREY
# blank rows to fill
for _ in range(4):
    blank = mtable.add_row().cells
    for i in range(3):
        blank[i].paragraphs[0].add_run("\n")
para(M.MATCH_BOUNDARY, italic=True, color=GREY, before=6)

# ---------- Section 6: Put Your Record to Work ----------
h_section(6, M.PUT_TO_WORK_TITLE)
para(M.PUT_TO_WORK_INTRO, after=6)
_ptw_lines = [2, 2, 3, 4]
_ptw_icons = ["resume", "about", "interview", "promotion"]
for _n, (title, purpose) in enumerate(M.PUT_TO_WORK_USES):
    h_block(title, color=GOLD, size=10.5, icon=_ptw_icons[_n])
    para(purpose, italic=True, color=GREY, size=9, after=2)
    for i in range(_ptw_lines[_n]):
        lp = doc.add_paragraph(); lp.paragraph_format.space_before = Pt(10); _bottom_border(lp, "C9C4B7", 4)

# ---------- Section 7 ----------
h_section(7, "Your index")
para("A running list of what you have, so any entry is a search away. Add a line whenever you add or revise an entry, and read it when a moment arrives and you need the right material fast.", after=6)
table = doc.add_table(rows=1, cols=len(M.INDEX_COLUMNS)); table.style = "Table Grid"
hdr = table.rows[0].cells
for i, c in enumerate(M.INDEX_COLUMNS):
    hdr[i].text = ""
    rp = hdr[i].paragraphs[0].add_run(c); rp.bold = True; rp.font.size = Pt(9); rp.font.color.rgb = NAVY
for _ in range(8):
    table.add_row()

# ---------- Section 8 ----------
h_section(8, "Optional maintenance checklists")
para("Two ways to keep the record current: capture when something worth keeping happens, and a light periodic sweep that catches what you missed. Neither is required, and neither is better; use whichever you will sustain, or both.", after=6)
h_block("Monthly sweep, about ten to fifteen minutes", color=NAVY, size=10.5, icon="monthly")
for m in M.MONTHLY:
    checkbox_line(m)
h_block("Quarterly review, about thirty minutes", color=NAVY, size=10.5, icon="quarterly")
for q in M.QUARTERLY:
    checkbox_line(q)
para(M.MAINT_NOTE, italic=True, color=GREY, before=6)

divider()
para("Tell me how it went: temidayoafonja.com/review.", bold=True, color=NAVY, before=4)

os.makedirs(os.path.dirname(OUT_DOCX), exist_ok=True)
doc.save(OUT_DOCX)
print("DOCX:", OUT_DOCX, os.path.getsize(OUT_DOCX), "bytes")

# ================= Markdown mirror =================
md = []
md.append(f"# {M.TITLE}")
md.append(f"*{M.SUBTITLE}*\n")
md.append(M.INTRO_LEAD + "\n")
for b, t in M.INTRO_BULLETS:
    md.append(f"- **{b}** {t}")
md.append("")
md.append(M.INTRO_SECTIONS + "\n")
md.append("*" + M.FORMATS_NOTE + "*\n")
md.append("---\n")

md.append("## Before you rebuild anything\n")
md.append(f"**{M.THREE_MOMENTS_LABEL}**  _{M.THREE_MOMENTS_HINT}_\n")
for _i in range(3):
    md.append("`__________________________________________________________________`")
md.append("")
md.append(f"**{M.WORDS_TITLE}.** _{M.WORDS_LEAD}_\n")
for w in M.WORDS_AFFIRMATIONS:
    md.append(f"- {w}")
md.append("")
md.append(f"**{M.WORDS_FILLIN}**")
md.append("`__________________________________________________________________`")
md.append("")
md.append("### Read It Back\n")
for _p in ["A sentence that shrinks my work",
           "What would have been different if I had not been there?",
           "The line I am learning to believe",
           "The entries that support it"]:
    _sep = "" if _p[-1:] in "?!.:" else "."
    md.append(f"**{_p}{_sep}**")
    md.append("`__________________________________________________________________`")
    md.append("`__________________________________________________________________`")
    md.append("")
md.append("---\n")

def md_field(label, hint, lines):
    h = f"  _{hint}_" if hint else ""
    sep = "" if label[-1:] in "?!.:" else "."
    md.append(f"**{label}{sep}**{h}")
    md.append("`__________________________________________________________________`")
    if lines > 1:
        md.append("`__________________________________________________________________`")
    md.append("")

md.append("## Section 1. Capture Log\n")
md.append("Catch work the moment it happens, in two minutes. Fill one block per event; expand only the strongest into Full Entries. Copy the block below for each new capture.\n")
for n in range(1,4):
    md.append(f"### Capture {n}\n")
    for label,hint,lines in M.CAPTURE_FIELDS: md_field(label,hint,lines)
    md.append("---\n")

md.append("## Section 2. Full Entries\n")
md.append("For work worth keeping in depth. Fill only the fields that apply. Copy the whole template for each new entry.\n")
for n in range(1,3):
    md.append(f"### Full Entry {n}\n")
    for ct,fields in M.FULL_ENTRY_CLUSTERS:
        md.append(f"**{ct}**\n")
        for label,hint,lines in fields: md_field(label,hint,lines)
    md.append(f"_{M.RECON_MARKER}_\n")
    md.append("---\n")

md.append("## Section 3. Corroboration list\n")
md.append("Who could confirm your work while they still remember it. Corroboration readiness, not networking.\n")
for n in range(1,5):
    md.append(f"### Corroboration {n}\n")
    for label,hint,lines in M.CORROB_FIELDS: md_field(label,hint,lines)
md.append(f"_{M.CORROB_NOTE}_\n")
md.append("---\n")

md.append("## Section 4. Translation and Proof Line workspace\n")
md.append("**Translation worksheet.** Work down the eight moves and apply the ones that fit, then the protection rule.\n")
for m in M.TRANSLATION_MOVES: md.append(f"- [ ] {m}")
md.append("")
md.append(f"_{M.TRANSLATION_PROTECTION}_\n")
md.append("**Proof Line builder.** Combine up to five ingredients into one portable sentence.\n")
for label,hint,lines in M.PROOFLINE_INGREDIENTS: md_field(label,hint,lines)
md.append("**Your Proof Line.**")
md.append("`__________________________________________________________________`\n")
md.append(f"_{M.PROOFLINE_RULE}_\n")
md.append("**Proof Lines kept.** A running list of your finished Proof Lines.\n")
for i in range(6): md.append("`__________________________________________________________________`")
md.append("\n---\n")

md.append(f"## Section 5. {M.MATCH_TITLE}\n")
md.append(M.MATCH_INTRO + "\n")
md.append("| " + " | ".join(M.MATCH_COLUMNS) + " |")
md.append("| " + " | ".join(["---"]*3) + " |")
md.append("| " + " | ".join(f"_{v}_" for v in M.MATCH_EXAMPLE) + " |")
for _ in range(4):
    md.append("| " + " | ".join([" "]*3) + " |")
md.append("")
md.append(f"_{M.MATCH_BOUNDARY}_\n")
md.append("---\n")

md.append(f"## Section 6. {M.PUT_TO_WORK_TITLE}\n")
md.append(M.PUT_TO_WORK_INTRO + "\n")
_ptw_md = [2, 2, 3, 4]
for _n,(title,purpose) in enumerate(M.PUT_TO_WORK_USES):
    md.append(f"### {title}\n")
    md.append(f"_{purpose}_\n")
    for _i in range(_ptw_md[_n]):
        md.append("`__________________________________________________________________`")
    md.append("")
md.append("---\n")

md.append("## Section 7. Your index\n")
md.append("A running list of what you have, so any entry is a search away.\n")
md.append("| " + " | ".join(M.INDEX_COLUMNS) + " |")
md.append("| " + " | ".join(["---"]*len(M.INDEX_COLUMNS)) + " |")
for _ in range(8):
    md.append("| " + " | ".join([" "]*len(M.INDEX_COLUMNS)) + " |")
md.append("\n---\n")

md.append("## Section 8. Optional maintenance checklists\n")
md.append("**Monthly sweep, about ten to fifteen minutes.**\n")
for m in M.MONTHLY: md.append(f"- [ ] {m}")
md.append("")
md.append("**Quarterly review, about thirty minutes.**\n")
for q in M.QUARTERLY: md.append(f"- [ ] {q}")
md.append("")
md.append(f"_{M.MAINT_NOTE}_")

os.makedirs(os.path.dirname(OUT_MD), exist_ok=True)
open(OUT_MD,"w",encoding="utf-8").write("\n".join(md))
print("MD (internal):", OUT_MD, os.path.getsize(OUT_MD), "bytes")
