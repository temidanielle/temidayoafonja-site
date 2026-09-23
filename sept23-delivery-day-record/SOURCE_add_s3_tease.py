# -*- coding: utf-8 -*-
"""Slide 3 only: let the boundary slide tease the seven categories of move.

The IT WILL row already promised "one category of move your evidence supports
testing". It never said there are seven, and it never said when. Both are the
reason to stay: the categories land at 38:00, long after the first read.

    before  Give you a defensible reading of your current position and one
            category of move your evidence supports testing.

    after   Give you a defensible reading of your current position, and at
            minute 38, all seven categories of move, then the one your
            evidence supports testing.

"all seven categories of move" is bold so the promise carries at a glance. It is
set as its own run inside the same text box, so the row keeps its font, size and
colour, the box does not move, and no shape is added.

Why this row and not a new line: the four rows end at 4.70in and the legal line
sits at 4.92in with the footer at 5.21in, so a fifth row would have to push
something. The tease also belongs in IT WILL on the merits, because it is a
statement of what the session does.

The boundary is unchanged. The row still promises ONE category at the end, not
seven recommendations, and "supports testing" is kept word for word.
"""
import copy, hashlib, os, shutil
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(
    HERE,
    "PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_60MIN_"
    "v2.0.7_CANDIDATE_s12-numbered_s20-reordered.pptx")
DST = os.path.join(
    HERE,
    "PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_60MIN_"
    "v2.0.7_CANDIDATE_s3-tease_s12-numbered_s20-reordered.pptx")

OLD = ("Give you a defensible reading of your current position and one category "
       "of move your evidence supports testing.")
PARTS = [
    ("Give you a defensible reading of your current position, and at minute 38, ", False),
    ("all seven categories of move", True),
    (", then the one your evidence supports testing.", False),
]
IT_WILL = 5          # shape index of the IT WILL value on slide 3


def build():
    shutil.copyfile(SRC, DST)
    p = Presentation(DST)
    sl = list(p.slides)
    assert len(sl) == 33, "slide count moved"
    s3 = sl[2]
    shp = list(s3.shapes)
    assert "What this read will and will not do" in shp[1].text_frame.text, \
        "slide 3 is not the boundary slide"
    assert shp[4].text_frame.text.strip() == "IT WILL", "the IT WILL row moved"

    para = shp[IT_WILL].text_frame.paragraphs[0]
    runs = [r for r in para.runs if r.text.strip()]
    assert len(runs) == 1 and runs[0].text.strip() == OLD, \
        f"IT WILL reads {shp[IT_WILL].text_frame.text!r}"

    base = runs[0]
    base.text = PARTS[0][0]
    prev = base
    for text, bold in PARTS[1:]:
        el = copy.deepcopy(base._r)
        prev._r.addnext(el)
        new = next(r for r in para.runs if r._r is el)
        new.text = text
        new.font.bold = bold
        prev = new
    base.font.bold = PARTS[0][1]

    p.save(DST)

    # ---- prove the pass was narrow --------------------------------------
    before, after = Presentation(SRC), Presentation(DST)
    bs, as_ = list(before.slides), list(after.slides)

    def face(x):
        return "\n".join(s.text_frame.text for s in x.shapes if s.has_text_frame)

    def note(x):
        return x.notes_slide.notes_text_frame.text if x.has_notes_slide else ""

    changed = [i + 1 for i in range(33)
               if face(bs[i]) != face(as_[i]) or note(bs[i]) != note(as_[i])]
    assert changed == [3], f"something changed outside slide 3: {changed}"

    old, new = list(bs[2].shapes), list(as_[2].shapes)
    assert len(old) == len(new), "slide 3 gained or lost a shape"
    for a, b in zip(old, new):
        assert (a.left, a.top, a.width, a.height) == (b.left, b.top, b.width, b.height), \
            f"slide 3 shape {a.name} moved"
        if a.has_text_frame and list(old).index(a) != IT_WILL:
            assert a.text_frame.text == b.text_frame.text, f"{a.name} text changed"

    got = new[IT_WILL].text_frame.text
    assert got == "".join(t for t, _ in PARTS), f"IT WILL reads {got!r}"
    assert "supports testing." in got, "the boundary phrase was lost"

    print("built", os.path.basename(DST))
    print("  changed:", changed)
    print("  IT WILL:", got)
    print("  sha256", hashlib.sha256(open(DST, "rb").read()).hexdigest())
    return DST


if __name__ == "__main__":
    build()
