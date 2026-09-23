# -*- coding: utf-8 -*-
"""Two changes to the presenting copy, on top of the slide 12 clarifier.

  1. DELETE the duplicate State Costs slide (slide 21). Slides 20 and 21 are
     identical, face and note. Removing 21 also realigns the hardcoded footer
     numbers, which were already written for a 33-slide deck: slide 22 already
     carries the number 21.

  2. REORDER the four rows on slide 20 so the block climbs rather than reading
     the matrix left to right:

         STAGNANT      neither axis accruing
         DEPTH TRAP    something real is being built, it does not travel
         FRAGILE       what you have travels, nothing is being added under it
         COMPOUNDING   both axes high

     Only the row CONTENT moves. Every chip, box and text frame keeps its exact
     position and size.

     The chip colour follows the meaning rather than the old zebra stripe: the
     first three rows take the light panel, and COMPOUNDING alone takes the navy
     panel with the gold label. That is the same relationship slide 12 already
     holds, where Compounding is the one inverted square. Under the old
     alternating stripe, reordering would have handed the boldest treatment to
     Stagnant, which is the opposite of what the sequence is saying.

     Nothing else on slide 20 changes: the eyebrow, title, subtitle, closing
     italic line, footer and the speaker note are all untouched. The note's
     "one pass, four states, no favourites" still holds, because it prohibits
     deepening a square for the room, not putting the four in an order.

The workbook was checked before reordering: pages 3 and 7 carry the 2x2 grid and
the placement tick list, but no state-costs sequence at all, so nothing in
participants' hands contradicts the new order.
"""
import copy, hashlib, os, shutil
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(
    HERE,
    "PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_60MIN_"
    "v2.0.7_CANDIDATE_slide12-clarifier.pptx")
DST = os.path.join(
    HERE,
    "PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_60MIN_"
    "v2.0.7_CANDIDATE_s12-clarifier_s20-reordered.pptx")

NAVY = RGBColor(0x0F, 0x23, 0x47)
PANEL = RGBColor(0xE9, 0xED, 0xF3)
GOLD = RGBColor(0xC9, 0xA8, 0x4C)

# shape indices on slide 20, one triple per row: (chip, label, cost)
ROWS = [(3, 4, 5), (6, 7, 8), (9, 10, 11), (12, 13, 14)]

# the new order, lowest to highest, with each state's own cost line carried over
ORDER = [
    ("STAGNANT",    "Neither axis is accruing."),
    ("DEPTH TRAP",  "Value is deep but context-bound."),
    ("FRAGILE",     "What travels today is no longer being renewed."),
    ("COMPOUNDING", "A strong current position still requires renewal."),
]


def carrier(shape):
    """The run that holds this shape's text, ignoring empty leading runs."""
    first = shape.text_frame.paragraphs[0]
    return next((r for r in first.runs if r.text.strip()), first.runs[0])


def build():
    shutil.copyfile(SRC, DST)
    p = Presentation(DST)
    sl = list(p.slides)
    assert len(sl) == 34, "slide count moved"

    # ---- 1. slide 20: reorder ------------------------------------------
    s20 = sl[19]
    shp = list(s20.shapes)
    assert "What each state costs" in shp[1].text_frame.text, "slide 20 moved"
    old_pairs = [(carrier(shp[l]).text, carrier(shp[c]).text) for _, l, c in ROWS]
    assert sorted(old_pairs) == sorted(ORDER), \
        f"slide 20 rows are not the four expected states: {old_pairs}"
    old_geom = [(s.left, s.top, s.width, s.height) for s in shp]

    for (chip_i, label_i, cost_i), (label, cost) in zip(ROWS, ORDER):
        last = label == "COMPOUNDING"
        chip = shp[chip_i]
        chip.fill.solid()
        chip.fill.fore_color.rgb = NAVY if last else PANEL
        lr = carrier(shp[label_i])
        lr.text = label
        lr.font.color.rgb = GOLD if last else NAVY
        carrier(shp[cost_i]).text = cost

    # ---- 2. delete the duplicate slide 21 -------------------------------
    s21 = sl[20]
    assert s21.shapes[1].text_frame.text == s20.shapes[1].text_frame.text, \
        "slide 21 is not a State Costs slide"
    lst = p.slides._sldIdLst
    entry = list(lst)[20]
    p.part.drop_rel(entry.rId)
    lst.remove(entry)

    p.save(DST)

    # ---- prove the pass was narrow --------------------------------------
    before, after = Presentation(SRC), Presentation(DST)
    bs, as_ = list(before.slides), list(after.slides)
    assert len(as_) == 33, f"expected 33 slides after deletion, got {len(as_)}"

    def face(x):
        return "\n".join(s.text_frame.text for s in x.shapes if s.has_text_frame)

    def note(x):
        return x.notes_slide.notes_text_frame.text if x.has_notes_slide else ""

    # after deletion, before[i] maps to after[i] up to 20, then shifts by one
    pairs = [(bs[i], as_[i]) for i in range(20)] + \
            [(bs[i + 1], as_[i]) for i in range(20, 33)]
    changed = [i + 1 for i, (b, a) in enumerate(pairs)
               if face(b) != face(a) or note(b) != note(a)]
    assert changed == [20], f"something changed outside slide 20: {changed}"

    a20 = as_[19]
    ashp = list(a20.shapes)
    assert len(ashp) == len(old_geom), "slide 20 gained or lost a shape"
    for (l, t, w, h), s in zip(old_geom, ashp):
        assert (s.left, s.top, s.width, s.height) == (l, t, w, h), \
            f"slide 20 shape {s.name} moved"
    new_pairs = [(carrier(ashp[l]).text, carrier(ashp[c]).text) for _, l, c in ROWS]
    assert new_pairs == ORDER, f"slide 20 order is wrong: {new_pairs}"
    assert note(a20) == note(bs[19]), "slide 20's note changed"

    # the slide 12 clarifier from the previous pass survived
    assert "Neither axis is moving." in face(as_[11]), "slide 12 clarifier lost"
    assert "It is just not being added to." in face(as_[11]), "slide 12 clarifier lost"

    print("built", os.path.basename(DST))
    print("  slides:", len(bs), "->", len(as_))
    print("  changed:", changed)
    print("  sha256", hashlib.sha256(open(DST, "rb").read()).hexdigest())
    return DST


if __name__ == "__main__":
    build()
