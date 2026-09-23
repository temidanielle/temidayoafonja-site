# -*- coding: utf-8 -*-
"""Slide 12 only: add the Fragile / Stagnant clarifier inside the two bottom
squares of the matrix.

One change, nothing else. Two new text lines are added, one inside the STAGNANT
square and one inside the FRAGILE square, sitting under the descriptor each
square already carries. They are placed inside the panels rather than beneath
the grid because the only clear space below the grid is the 0.24in gap between
the Optionality axis caption and the footer, which will not hold two lines
without moving the caption or the footer.

The label prefixes from the spoken version ("FRAGILE:", "STAGNANT:") are not
repeated in the text: each line sits directly under the label it belongs to,
inside the same square, so the prefix is carried visually.

Each line is cloned from the descriptor above it, so it inherits the deck's own
DM Sans / 5A6B84 styling rather than being drawn with guessed values, and is
then stepped down from 11pt to 9.5pt so it reads as secondary furniture and not
as a second descriptor.

Nothing else on slide 12 moves. No other slide, face or note, is touched.
"""
import copy, hashlib, os, shutil
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "src.pptx")
DST = os.path.join(
    HERE,
    "PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_60MIN_"
    "v2.0.7_CANDIDATE_slide12-clarifier.pptx")

# (index of the descriptor to clone, new text)
ADD = [
    (11, "Neither axis is moving."),
    (14, "What you built travels. It is just not being added to."),
]
# Both lines share a top edge so the bottom row stays level. The bottom panels
# run 3.28 to 4.60in; the descriptors end near 4.05in. 4.18in leaves the
# descriptor its air and keeps two wrapped lines clear of the panel edge.
TOP = 4.18
HEIGHT = 0.40
SIZE = Pt(9.5)


def set_text(shape, text):
    """Rewrite the shape's text in place, keeping the carrier run's formatting."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    keep = next((r for r in first.runs if r.text.strip()), first.runs[0])
    for p in tf.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    for r in list(first.runs):
        if r._r is not keep._r:
            r._r.getparent().remove(r._r)
    keep.text = text
    return keep


def next_shape_id(slide):
    ids = [int(el.get("id")) for el in slide.shapes._spTree.iter()
           if el.tag.endswith("}cNvPr") and el.get("id")]
    return max(ids) + 1


def build():
    shutil.copyfile(SRC, DST)
    p = Presentation(DST)
    sl = list(p.slides)
    s12 = sl[11]
    shp = list(s12.shapes)

    assert "The four states" in shp[1].text_frame.text, "slide 12 is not the matrix"
    assert shp[10].text_frame.text.strip() == "STAGNANT"
    assert shp[13].text_frame.text.strip() == "FRAGILE"
    assert shp[9].top == Inches(3.28) and shp[9].height == Inches(1.32)

    sid = next_shape_id(s12)
    for src_idx, text in ADD:
        source = shp[src_idx]
        el = copy.deepcopy(source._element)
        s12.shapes._spTree.append(el)
        new = s12.shapes[-1]
        for cn in new._element.iter():
            if cn.tag.endswith("}cNvPr"):
                cn.set("id", str(sid))
                cn.set("name", "Clarifier;%d;p12" % sid)
                sid += 1
        new.left, new.width = source.left, source.width
        new.top, new.height = Inches(TOP), Inches(HEIGHT)
        run = set_text(new, text)
        run.font.size = SIZE

    p.save(DST)

    # ---- prove the pass was narrow -------------------------------------
    before, after = Presentation(SRC), Presentation(DST)
    bs, as_ = list(before.slides), list(after.slides)
    assert len(bs) == len(as_) == 34, "slide count moved"

    def faces(deck):
        return ["\n".join(s.text_frame.text for s in x.shapes if s.has_text_frame)
                for x in deck]

    def notes(deck):
        return [x.notes_slide.notes_text_frame.text if x.has_notes_slide else ""
                for x in deck]

    bf, af, bn, an = faces(bs), faces(as_), notes(bs), notes(as_)
    changed_faces = [i + 1 for i in range(34) if bf[i] != af[i]]
    changed_notes = [i + 1 for i in range(34) if bn[i] != an[i]]
    assert changed_faces == [12], f"a face changed outside slide 12: {changed_faces}"
    assert changed_notes == [], f"a note changed: {changed_notes}"
    assert af[11] == bf[11] + "\n" + ADD[0][1] + "\n" + ADD[1][1], \
        "slide 12 gained something other than the two lines"

    # every pre-existing shape on slide 12 keeps its exact geometry
    old, new = list(bs[11].shapes), list(as_[11].shapes)
    assert len(new) == len(old) + 2, "slide 12 shape count is wrong"
    for a, b in zip(old, new):
        assert (a.left, a.top, a.width, a.height) == (b.left, b.top, b.width, b.height), \
            f"shape {a.name} moved"
        if a.has_text_frame:
            assert a.text_frame.text == b.text_frame.text, f"shape {a.name} text changed"

    print("built", os.path.basename(DST))
    print("  faces changed:", changed_faces, " notes changed:", changed_notes)
    print("  sha256", hashlib.sha256(open(DST, "rb").read()).hexdigest())
    return DST


if __name__ == "__main__":
    build()
