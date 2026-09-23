# -*- coding: utf-8 -*-
"""Slide 12 only: number the four squares so the climb is visible.

    1  STAGNANT      2  DEPTH TRAP      3  FRAGILE      4  COMPOUNDING

The numeral is added as its own run in front of each existing label run, so the
label keeps its exact font, size, weight and colour and nothing on the slide
moves. No new shapes, no geometry change.

Colour: the numeral always takes the panel's secondary accent, so it reads as an
order marker rather than as part of the state's name.

    light panels   label is navy 0F2347   ->  numeral gold C9A84C
    navy panel     label is gold C9A84C   ->  numeral B8C5D9, the tone that
                                              panel already uses for its text

The numbers match the order slide 20 now runs in, so the two slides agree.
"""
import copy, hashlib, os, shutil
from pptx import Presentation
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(
    HERE,
    "PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_60MIN_"
    "v2.0.7_CANDIDATE_s12-clarifier_s20-reordered.pptx")
DST = os.path.join(
    HERE,
    "PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_60MIN_"
    "v2.0.7_CANDIDATE_s12-numbered_s20-reordered.pptx")

GOLD = RGBColor(0xC9, 0xA8, 0x4C)
ONNAVY = RGBColor(0xB8, 0xC5, 0xD9)

# shape index on slide 12 -> (expected label, number, numeral colour)
LABELS = {
    10: ("STAGNANT", "1", GOLD),
    4:  ("DEPTH TRAP", "2", GOLD),
    13: ("FRAGILE", "3", GOLD),
    7:  ("COMPOUNDING", "4", ONNAVY),
}
GAP = "   "     # three spaces; Montserrat's space is narrow


def carrier(shape):
    first = shape.text_frame.paragraphs[0]
    return next((r for r in first.runs if r.text.strip()), first.runs[0])


def build():
    shutil.copyfile(SRC, DST)
    p = Presentation(DST)
    sl = list(p.slides)
    assert len(sl) == 33, "slide count moved"
    s12 = sl[11]
    shp = list(s12.shapes)
    assert "The four states" in shp[1].text_frame.text, "slide 12 moved"

    for idx, (expect, num, colour) in LABELS.items():
        label = carrier(shp[idx])
        assert label.text.strip() == expect, \
            f"shape {idx} is {label.text.strip()!r}, expected {expect!r}"
        assert not label.text.strip()[0].isdigit(), "already numbered"
        el = copy.deepcopy(label._r)
        label._r.addprevious(el)
        new = shp[idx].text_frame.paragraphs[0].runs[
            [r._r for r in shp[idx].text_frame.paragraphs[0].runs].index(el)]
        new.text = num + GAP
        new.font.color.rgb = colour

    p.save(DST)

    # ---- prove the pass was narrow --------------------------------------
    before, after = Presentation(SRC), Presentation(DST)
    bs, as_ = list(before.slides), list(after.slides)

    def face(x):
        return "\n".join(s.text_frame.text for s in x.shapes if s.has_text_frame)

    def note(x):
        return x.notes_slide.notes_text_frame.text if x.has_notes_slide else ""

    changed = [i + 1 for i in range(33)
               if face(bs[i]) != face(as_[i]) or note(bs[i]) != note(as_[i])]
    assert changed == [12], f"something changed outside slide 12: {changed}"

    old, new = list(bs[11].shapes), list(as_[11].shapes)
    assert len(old) == len(new), "slide 12 gained or lost a shape"
    for a, b in zip(old, new):
        assert (a.left, a.top, a.width, a.height) == (b.left, b.top, b.width, b.height), \
            f"slide 12 shape {a.name} moved"

    for idx, (expect, num, _) in LABELS.items():
        got = new[idx].text_frame.text
        assert got == num + GAP + expect, f"label {idx} reads {got!r}"
    # every other shape's text is untouched, including the clarifier lines
    for i, (a, b) in enumerate(zip(old, new)):
        if i in LABELS or not a.has_text_frame:
            continue
        assert a.text_frame.text == b.text_frame.text, f"shape {a.name} text changed"

    print("built", os.path.basename(DST))
    print("  changed:", changed)
    print("  labels:", [new[i].text_frame.text for i in (10, 4, 13, 7)])
    print("  sha256", hashlib.sha256(open(DST, "rb").read()).hexdigest())
    return DST


if __name__ == "__main__":
    build()
