# -*- coding: utf-8 -*-
"""Free Flagship SOP — v2.0.8 CANDIDATE -> v2.0.9 CANDIDATE.

Three operational changes, no methodology change:

  POLLS       The blanket ban becomes a precise rule. Nothing diagnostic,
              score-revealing, state-revealing, employer-revealing or
              decision-revealing runs inside the recorded assessment. One
              optional anonymous arrival poll may run BEFORE the recording
              starts, while the dated holding slide is up, and only if the
              platform makes it easy.

  COMMERCIAL  The retired Private Capability Position Read is removed as an
              offer and replaced by the current individual architecture: Career
              Evidence Starter free, Keep the Proof $49, Field Kit $150, Career
              Move Review $500 with qualification before payment. Career Move
              Review is NOT on the participant-facing continuation slide,
              because no route for it exists in the repository. Its route
              verification is recorded as an outstanding launch condition.

  BOUNDARY    Q&A gains the destination-portability boundary: this session did
              not read anyone's destination, and generic portability is not
              destination-specific portability.

Everything else stands, including the master reconciliation table, which needs
no timing edit: the v2.0.6 opening moves ten seconds between slides 1 and 2 and
leaves every block boundary exactly where it was.
"""
import hashlib, os, shutil
import docx

SRC = ("sept23-v205-assets/"
       "Capability_Formation_Free_Flagship_SOP_Sept23_2026_60MIN_v2.0.8_CANDIDATE.docx")
OUT = "scratchpad/sept23/out"
DST = (f"{OUT}/"
       "Capability_Formation_Free_Flagship_SOP_Sept23_2026_60MIN_v2.0.9_CANDIDATE.docx")
STAMP = "Saturday, September 19, 2026 at 8:30 AM CT"
OLD_STAMP = "Friday, September 4, 2026 at 6:40 AM CT"

# Straight substitutions, applied run by run so formatting survives.
SUBS = [
    ("v2.0.8", "v2.0.9"),
    ("v2.0.5", "v2.0.6"),
    (OLD_STAMP, STAMP),
]

# Located edits. Each is (paragraph index, expected fragment, replacement).
TRIMS = [
 (113, "Two categories, previously listed as one.",
  "Two categories. The first is live and outside this document's scope; the second "
  "does not exist yet. Changing anything in the first is a separate decision: this "
  "SOP points at those assets, it neither governs nor republishes them."),
 (117, "Field Kit, reached through the branded route",
  "Field Kit, reached through temidayoafonja.com/fieldkit, which is the deck face "
  "route, the QR target and the chat message. It is the evergreen destination "
  "because a recorded replay outlives any storefront URL."),
 (118, "Current website copy, including the page",
  "Current website copy. The live site is the source of truth for pricing."),
]

PARA = [
 (37,
  "No poll, no show of hands, no room-level state distribution, no anonymous count.",
  "NO DIAGNOSTIC POLL INSIDE THE RECORDED ASSESSMENT. Nothing that asks anyone to "
  "reveal a score, a state, an employer, a circumstance or a decision. No show of "
  "hands, no room-level state distribution, no anonymous count of squares. The "
  "optional “What moved” interaction is in the appendix and is OUT of the timed "
  "core. One optional anonymous arrival poll is permitted, and it runs BEFORE the "
  "recording starts: see below."),
 (50,
  "Do NOT offer the Private Capability Position Read",
  "If a question is really a request for a personal decision, name the boundary "
  "once and stop. If it asks whether experience will transfer to a named role, "
  "function, industry or employer, say plainly that this session did not read that "
  "destination: generic portability is not destination-specific portability. Do NOT "
  "sell the Career Move Review from the stage while its route is unverified, and do "
  "not mention the retired Private Capability Position Read at all."),
 (62,
  "No pricing, no Field Kit pitch, no Private Read pitch, no product CTA.",
  "No pricing, no Field Kit pitch, no Career Move Review pitch, no product CTA."),
 (75,
  "Confirm the Private Capability Position Read remains off",
  "Verify the current Field Kit route and price immediately before delivery. "
  "Confirm the Career Move Review is still absent from the participant-facing "
  "continuation sequence: it stays off until its fulfillment route is verified "
  "operational."),
 (121,
  "Private Capability Position Read materials, booking flow and fulfillment.",
  "Career Move Review fulfillment route. The offer and price are decided; no route "
  "for it exists in the repository or on the live site, so it is off the "
  "continuation slide in deck v2.0.6 and out of Q&A until a verified individual "
  "route can be named here. The Private Capability Position Read is RETIRED."),
]

# Section 4 gains the arrival poll, placed immediately after the privacy bullets
# so it reads as the one exception to the rule above it rather than as a new
# idea somewhere else in the document.
ARRIVAL_POLL = (
 "OPTIONAL ANONYMOUS ARRIVAL POLL, PRE-RECORDING ONLY. While the dated holding "
 "slide is up and the recording is OFF, one anonymous poll may ask “What brought "
 "you here?” with these options: my work feels static; I am considering an "
 "internal move; I am considering a career pivot; my role has changed around me; "
 "I want to know what would still count somewhere else. Close it before the "
 "recording starts. Do not read individual answers aloud and do not report a "
 "distribution into the recorded core. If the platform does not make it easy, skip "
 "it. Nothing depends on it.")

# The offer table, rebuilt to the current individual architecture.
OFFERS = [
 ("Offer", "Position"),
 ("September 23 guided assessment", "$0. Free for this launch."),
 ("Career Evidence Starter", "Free. The entry artifact. Not pitched from the stage."),
 ("Keep the Proof",
  "$49. For someone who has done the work but has no reliable record of it. A "
  "different problem from this session's, so it is NOT in the recorded continuation "
  "sequence. It belongs in follow-up, matched to that need."),
 ("Capability Formation Field Kit",
  "Exactly $150. Private, self-guided, repeatable. Not discounted, not bundled, "
  "never free. The only paid route named in the recorded sequence, and its price is "
  "spoken live rather than printed on the slide."),
 ("Career Move Review — ROUTE NOT VERIFIED",
  "$500, one public price, qualification before payment. The current advisory offer "
  "for someone with a named destination. NOT on the continuation slide and NOT sold "
  "from the stage until a verified individual route exists. Never route it to the "
  "enterprise Advisory page and never reuse a retired Private Read link."),
 ("Private Capability Position Read — RETIRED",
  "Withdrawn. Not priced, not pitched, not booked, and it does not return under "
  "that name. Career Move Review replaces it."),
 ("Former paid group workshop",
  "RETIRED from active public sale. May remain archived as an internal asset. Must "
  "NOT be presented as the current public next step."),
]


def iter_runs(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            yield r
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        yield r


def set_para(para, text):
    runs = para.runs
    keep = next((r for r in runs if r.text.strip()), runs[0])
    keep.text = text
    for r in runs:
        if r._r is not keep._r:
            r._r.getparent().remove(r._r)


def set_cell(cell, text):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    set_para(p, text)


def build():
    os.makedirs(OUT, exist_ok=True)
    shutil.copyfile(SRC, DST)
    d = docx.Document(DST)

    counts = {a: 0 for a, _ in SUBS}
    for r in iter_runs(d):
        t = r.text
        for a, b in SUBS:
            if a in t:
                counts[a] += t.count(a)
                t = t.replace(a, b)
        if t != r.text:
            r.text = t
    missing = [a for a, n in counts.items() if n == 0]
    assert not missing, f"substitution never matched: {missing}"

    for idx, expect, new in PARA + TRIMS:
        p = d.paragraphs[idx]
        assert expect in p.text, f"P{idx}: expected {expect[:40]!r}, found {p.text[:70]!r}"
        set_para(p, new)

    # the arrival poll, as a new bullet cloned from the bullet above it
    import copy
    anchor = d.paragraphs[37]
    el = copy.deepcopy(anchor._element)
    anchor._element.addnext(el)
    from docx.text.paragraph import Paragraph
    set_para(Paragraph(el, anchor._parent), ARRIVAL_POLL)

    # the offer table
    t = d.tables[4]
    assert "Private Capability Position Read" in t.rows[3].cells[0].text, \
        "offer table moved"
    template = copy.deepcopy(t.rows[1]._element)
    while len(t.rows) > 1:
        t._element.remove(t.rows[-1]._element)
    for ci, val in enumerate(OFFERS[0]):
        set_cell(t.rows[0].cells[ci], val)
    for row in OFFERS[1:]:
        t._element.append(copy.deepcopy(template))
        for ci, val in enumerate(row):
            set_cell(t.rows[-1].cells[ci], val)

    d.save(DST)
    return DST


if __name__ == "__main__":
    path = build()
    print("built", os.path.basename(path))
    print("  sha256", hashlib.sha256(open(path, "rb").read()).hexdigest())
