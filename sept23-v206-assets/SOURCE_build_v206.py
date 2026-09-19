# -*- coding: utf-8 -*-
"""Stay or Leave? 60-minute flagship — v2.0.5 CANDIDATE -> v2.0.6 CANDIDATE.

A focused revision of an approved deck, not a rebuild. The learning spine is
untouched: question, initial read, evidence, corrected read, interpretation,
move category, Next-Move Note.

WHAT CHANGES

  OPENING      The approved audit-to-cybersecurity warrant moves out of slide
               5's optional block and into the core spoken opening on slide 2,
               where it arrives before any terminology. Nothing is invented: the
               story ships with the words already approved in this package.

  LANGUAGE     Slides 4 and 5 ask the ordinary question first and name the term
               second. Density and Optionality both stay. Optionality's
               explanation is rewritten so it no longer reads as a claim about
               employer recognition.

  SLIDE 18     Simplified to three tiers, a boundary band and the dual-axis
               rule. The procedure it used to print in full already lives on
               workbook page 4 and in the presenter note, so nothing is lost.

  PORTABILITY  The four-question audit appears ONCE, on slide 22, as the
               interpretive frame for the Next-Move Note.

  CATEGORIES   Three of the seven one-line readings are rewritten. The seven
               names, and the rule that no square prescribes a category, stand.

  CLOSE        Slide 25 returns to the question the session opened with.

  COMMERCIAL   The retired Private Capability Position Read is replaced by the
               current architecture in the speaker notes. Career Move Review is
               NOT placed on the participant-facing continuation slide, because
               no route for it exists anywhere in the repository.

WHAT DOES NOT CHANGE

  The 12 statements, the 1-5 scale, the 90-day window, both /30 totals, the
  evidence protocol, the 3-versus-? distinction, the boundary and sensitivity
  methodology, the four states, the seven category names, private scoring, the
  Next-Move Note, the recording boundary at 50:00, the hidden-slide set, and
  every timing cue from 5:00 onward. The workbook is not touched at all.

TIMING. Two cues move, both inside the opening, to seat a 30-45 second lived
warrant without lengthening the session: slide 1 gives up ten seconds and slide 2
takes them. Slides 3, 4 and 5 keep the cues they had, the 2:00 and 5:00 block
boundaries are exact, every cue from 5:00 onward is untouched, and the SOP's
master reconciliation table needs no timing edit at all.
"""
import copy, hashlib, os, re, shutil
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor

SRC = ("sept23-v205-assets/"
       "PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_60MIN_"
       "v2.0.5_CANDIDATE.pptx")
OUT = "scratchpad/sept23/out"
DST = (f"{OUT}/PRESENTER_VERSION_Stay_or_Leave_Live_Career_Growth_Assessment_"
       "60MIN_v2.0.6_CANDIDATE.pptx")

NAVY = RGBColor(0x0F, 0x23, 0x47)

# ── slide faces ─────────────────────────────────────────────────────────────
# Each entry is (shape index, new text). Shape indices are asserted against the
# text they are replacing, so a reordered deck fails the build rather than
# writing the right words onto the wrong box.
FACE_EDITS = {
 4: [
   (1, "Density", "Is the work still building you?"),
   (2, "Learning velocity. Whether this environment is still forming you.",
       "This reading calls that Density."),
   (3, "Density is whether the environment is still building capability in you.",
       "Density reads whether this environment is still forming you. It reads your "
       "conditions, not your effort. A role can be demanding, important and "
       "exhausting and still stop building you."),
   (5, "Density is not motivation, hours, pressure or effort.",
       "Not how busy you are. Not how much pressure you carry. Not how much the "
       "organization depends on you. The question is whether the work is still "
       "asking more of your judgment than it did ninety days ago."),
 ],
 5: [
   # Kept to one line at the deck's 34pt title size. The longer phrasing
   # "Would what you are building still travel?" wrapped to two lines and
   # collided with the subtitle beneath it.
   (1, "Optionality", "Will what you build still travel?"),
   (2, "Career portability. Whether what you built travels.",
       "This reading calls that Optionality."),
   (3, "Optionality is whether the capability you are building would be recognised",
       "Optionality reads whether what this work is building stays useful beyond "
       "your current context, and whether you could make it legible to someone "
       "outside your organization using evidence rather than your title."),
   (5, "A low Optionality score may reflect context-bound capability",
   # Three lines, matching the panel on slide 4. The causes of a low score are
   # taught at slide 15 and in the appendix, so the panel carries the boundary
   # that matters most here: legibility is not recognition.
       "A low score is not a verdict on your ability. And translation is not "
       "recognition. Making your work legible removes a barrier. It does not decide "
       "what another organization does with what it sees."),
 ],
 18: [
   (2, "The neighbouring-score sensitivity check.",
       "The sensitivity check. An operational convention, not a statistical "
       "confidence interval."),
   (5, "No unsupported-item sensitivity range is required.",
       "Use the corrected total normally."),
   (6, "Use the corrected total normally, subject to the boundary band.",
       "No sensitivity range is required."),
   (10, "PROVISIONAL  —  RUN THE NEIGHBOURING-SCORE SENSITIVITY CHECK",
        "PROVISIONAL.  RUN THE NEIGHBORING-SCORE SENSITIVITY CHECK."),
   (12, "Total the axis using your numbers exactly as written.",
        "If the range touches 17 to 21, hold the axis as boundary."),
   (16, "Bounded inside the 1 to 5 scale",
        "17 TO 21 IS THE BOUNDARY BAND. A boundary reading overrides the normal "
        "high and low split."),
   (22, "INCOMPLETE — do not place on that axis today.",
        "INCOMPLETE. Do not place on that axis today. Incomplete is a legitimate "
        "reading, not a failure."),
   (23, "ONE axis in the 17–21 band",
        "ONE axis in the band: read the two adjacent states the other axis selects.\n"
        "BOTH axes in the band: record Boundary and assign no single state."),
 ],
 20: [
   (15, "This is the short reading. The fuller cost of your own square",
        "That is the reading this assessment gives you today, and it is complete. "
        "Working a square in depth, against a named destination, and rerunning it as "
        "conditions change is what continued work adds."),
 ],
 21: [
   # All three stay on one line. The column is 5.48in at 10.6pt, and a
   # two-line reading breaks the row rhythm the slide was laid out for.
   (10, "The capability exists. The language for it does not.",
        "The capability exists. The evidence for it is not yet legible outside."),
   (14, "The work is good and the wrong people have seen it.",
        "Reach people who can evaluate or use it. Visibility is not evidence."),
   (22, "Change the work, not the employer, first.",
        "Test whether changing the work or conditions could repair the problem."),
 ],
 25: [
   # Two lines of about thirty characters, which is what the 31pt title box
   # holds. The longer phrasing wrapped to four lines and painted over the
   # eyebrow above and the gold rule below.
   (1, "You have read your position.",
       "You came in asking stay or leave.\nYou leave with a clearer read."),
   (3, "Put a rescore date ninety days out",
       "I have not made that decision for you, and I should not. What you have is a "
       "clearer read of what this work is building, what may travel, where the "
       "evidence is thin, and what deserves testing next. Put a date ninety days "
       "from now on your note, then come back and see what changed."),
 ],
}

# Slide 18: the four procedure lines and their bullet dots. The steps they print
# are on workbook page 4 and in the presenter note; printing them on the slide
# as well is what made this the densest participant-facing slide in the deck.
S18_DELETE = [13, 14, 15, 17, 18]
S18_GEOM = {                       # shape index: (top, left, width, height) in inches
    7: (2.50, 0.62, 0.74, 0.86), 8: (2.50, 0.62, 0.74, 0.86),
    9: (2.50, 1.54, 7.84, 0.86), 10: (2.62, 1.76, 7.40, 0.22),
    11: (2.98, 1.76, 0.09, 0.09), 12: (2.90, 1.98, 7.18, 0.30),
    19: (3.52, 0.62, 0.74, 0.50), 20: (3.52, 0.62, 0.74, 0.50),
    21: (3.52, 1.54, 7.84, 0.50), 22: (3.52, 1.76, 7.40, 0.50),
    16: (4.18, 0.62, 8.76, 0.28),
    23: (4.58, 0.62, 8.76, 0.50),
}
# Slide 25's body needs another line of height for the longer close.
S25_GEOM = {3: (3.02, 0.62, 6.90, 1.20)}
# Slide 22's subtitle box carries the portability frame as a second line.
S22_GEOM = {2: (1.40, 0.62, 8.76, 0.58)}

PORTABILITY_FRAME = ("What travels?   ·   What does not?   ·   What can I prove?   "
                     "·   What must I relearn?")

# ── speaker notes ───────────────────────────────────────────────────────────
NOTES = {}

NOTES[1] = """START THE RECORDING ON THIS SLIDE. This is the first evergreen teaching slide. The dated holding slide is shown BEFORE recording begins and is never recorded.

TIMING: 0:00-0:30. THIRTY SECONDS.

FACILITATION: welcome, name the session, name the method, state the duration. Thirty seconds, and biography is capped at TWO SENTENCES. The lived example that earns the method comes on the next slide, so do not start it here.

WHAT TO SAY: "This is a live guided assessment. In the next fifty minutes you will read your own position twice, once from memory and once against evidence, and leave with a note about what your evidence supports testing next. The last ten minutes are live questions and are not recorded."

EVERGREEN-SAFE: no date on this slide by design."""

NOTES[2] = """RECORDING ON.

TIMING: 0:30-1:20. FIFTY SECONDS. The reframe is the first fifteen. The rest is the story below.

FACILITATION: read both boxes, then go straight into the story and let the story do the landing. Fifteen seconds on the reframe, the rest on the warrant. Recognition before terminology: the room needs to have heard one real example of what travels and what does not before Density and Optionality arrive as words.

SPOKEN WARRANT - SAY IT, DO NOT SKIP IT. Approximately thirty to forty-five seconds, and no longer:
"When I moved from audit into cybersecurity, I did not start from zero. Evidence, controls and risk travelled with me. Cybersecurity language, privacy frameworks and technical context did not. Portability was not plug-and-play. Something important came with me, and the new context still required real relearning."

Then land the point in one line: I did not start from zero, and I did not carry everything either.

DO NOT ADD TO THE STORY. No dates, no employers, no metrics, no reactions, no reason for the move, no outcome, and nothing further about what transferred. This is warrant, not biography. If you find yourself explaining the move, you have already gone too long.

Do not answer the stay/leave question. Say plainly that the session reads the position underneath it."""

NOTES[3] = """RECORDING ON. THIS IS THE RECORDING AND PRIVACY DISCLOSURE. It must be delivered inside the first five minutes. Deliver it here, at minute 1:20.

TIMING: 1:20-2:00. FORTY SECONDS.

FACILITATION: deliver the disclosure, name the boundary, point at the workbook, move. Forty seconds total. Do not let this become a preamble.

SAY IT ONCE, CALMLY, AND DO NOT EXPAND IT INTO A DISCLAIMER: "The teaching portion is recorded so people can revisit the method. Your scores and your private workbook stay yours. The live Q&A at the end is not part of the evergreen replay."

WORKBOOK: page 1 carries the same boundary and privacy language. Tell the room the workbook is open in full from the start. Nothing is released in stages.

Participant microphones are muted by default for the whole recorded core."""

NOTES[4] = """RECORDING ON.

TIMING: 2:00-3:30. NINETY SECONDS.

FACILITATION: ask the question on the slide first and let it sit for a beat. Only then give it the name. The room recognizes the question before it learns the word.

WHAT TO SAY, in your own words: is the work still asking more of your judgment than it used to? That is what Density reads. It is about your conditions, not your effort, which is why a role can be demanding and important and still stop forming you.

Give the one correction that matters, that busy is not the same as forming, and stop. The full teaching treatment is downstream, not here.

Do not take questions on the definition now. Say they will make more sense once people have scored."""

NOTES[5] = """RECORDING ON.

TIMING: 3:30-5:00. NINETY SECONDS.

FACILITATION: same shape as the last slide. Ask the question, let it sit, then name it. The story on slide 2 was the warrant for this slide, so you can refer back to it in a phrase rather than retelling it.

WHAT TO SAY, in your own words: would what you are building still be useful somewhere else, and could you make it legible to someone outside your organization using evidence rather than your title? That is Optionality.

SAY THE BOUNDARY AS WRITTEN. A low score is not a verdict on ability, and translation is not recognition. Making your capability legible removes a barrier. It does not decide what another organization does with what it sees, and this session does not predict hiring, promotion, sponsorship or pay.

IF THE ROOM WANTS THE CAUSES, they are taught at slide 15 and in the appendix: a low score may mean the capability is context-bound, that the evidence is hard to read from outside, or that the right people have not seen the work. Do not rank them and do not imply one is more common than another. This session has no measurement that supports a prevalence claim.

Next slide begins the private scoring. Tell the room to have the workbook open at page 2."""

NOTES[16] = """RECORDING ON. PRIVATE WRITING.  WORKBOOK PAGES 5 AND 6.

TIMING: 19:00-31:00. TWELVE MINUTES. PROTECTED AND NON-NEGOTIABLE.

FACILITATION: read the instruction once, then hold this slide and stay quiet. One short orientation cue at the midpoint, then the two warnings. Nothing else.

NEVER RECOVER LOST TIME HERE. This block is the evidentiary centre of the session and the thing that makes it an assessment rather than a talk. If you are behind at minute 19, take the time from the 45:00-48:00 continuation sequence or the 48:00-50:00 buffer, never from this block.

This slide holds for the full twelve minutes. Do not advance and do not fill the silence.

ONE MIDPOINT CUE AT ABOUT 25:00. This is facilitator presence, not teaching. Say it once, quietly, then return to silence: "If you are around statement six, you are on pace. Keep the evidence phrases short." Do not add anything to it, do not answer a question in it, and do not restate the protocol.

Give a two-minute warning at 29:00 and a thirty-second warning at 30:30. Keep your own microphone quiet the rest of the time.

Short evidence phrases, not essays. Twelve statements in twelve minutes is sixty seconds each."""

NOTES[18] = """RECORDING ON.  WORKBOOK PAGE 7.

TIMING: 32:30-34:00. NINETY SECONDS.

FACILITATION: thirty seconds per tier. Teach all three even though most of the room will only need one.

THE SLIDE IS NOW THE SUMMARY, NOT THE SPECIFICATION. The full procedure is on workbook page 4 and is repeated below. Read the three tiers off the slide and let the workbook carry the mechanics.

Teach all three tiers explicitly. A participant with no marked items can sit through the procedure believing they must run it. Say plainly that they do not.

NONE: no range required. 1 OR 2: the axis is provisional, run the check. 3 OR MORE: the axis is incomplete and you do not place on it today.

THE FULL SENSITIVITY PROCEDURE, for the room that asks and for workbook page 4: total the axis using the numbers exactly as written; total it again with every marked item one point lower; total it again one point higher. Every neighbouring score stays bounded inside the 1 to 5 scale, so a 1? is never tested at 0 and a 5? never at 6. If the range includes any total from 17 to 21, do not force a state, and hold the axis as boundary.

THE DUAL-AXIS BOUNDARY RULE, said plainly:
If exactly ONE axis is inside 17 to 21, read the two adjacent states that the OTHER axis selects. A boundary Density with high Optionality reads across Compounding and Fragile.
If BOTH axes are inside 17 to 21, record Boundary and assign no single state at all.

The governing instrument (v5.3.1) states that when any total touches the band, no single state is assigned and the placement is provisional. The one-axis case follows directly. The instrument does not address the both-axes case, where four squares would qualify rather than two; recording Boundary is the reading consistent with "no single state is assigned". This is documented in the change log.

Say the boundary out loud: this is an operational convention, not a statistical confidence interval, and it carries no margin of error. Boundary and incomplete are legitimate outcomes.

METHOD UNCHANGED IN v2.0.6. Only the slide face was simplified. The tiers, the band, the bounded neighbouring scores and the dual-axis rule are exactly what v2.0.5 carried."""

NOTES[19] = """RECORDING ON. PRIVATE WRITING.  WORKBOOK PAGE 7.

TIMING: 34:00-35:00. SIXTY SECONDS.

FACILITATION: give the instruction, then let them place. Do not narrate over the writing.

State plainly that boundary and incomplete are real readings, not failures.

NO DIAGNOSTIC POLL INSIDE THE RECORDED ASSESSMENT. No staged reveal, no room-level state distribution, no anonymous count, and nothing that asks anyone to reveal a score, a state, an employer or a decision. Scores, states, employers and circumstances stay private. The optional "what moved" interaction is in the appendix and is OUT of the timed core.

The one poll this session permits is the optional anonymous arrival poll, which runs BEFORE the recording starts while the dated holding slide is up. It is specified in the SOP. It never runs here.

Do not invite anyone to name their square, their confidence number, their employer or their circumstances."""

NOTES[20] = """RECORDING ON.

TIMING: 35:00-38:00. THREE MINUTES. Forty-five seconds per state.

FACILITATION: one pass, four states, no favourites. Keep your eye on the clock, this block overruns easily.

Give each state its essential reading and stop. Do NOT teach the heaviest square for this room and do NOT deepen a state based on audience distribution. Those are Intensive Edition mechanics and there is no room for them in sixty minutes.

TODAY IS COMPLETE. The closing line on this slide says the reading they have is the one this assessment gives, and that it is complete. Say it that way. Do not say the fuller cost is held back, and do not imply the real interpretation sits behind a paid product. What continued work adds is depth, repetition, a named destination and private use, not a withheld answer.

If you are running behind, this is NOT the block to cut. Take it from 45:00-48:00."""

NOTES[21] = """RECORDING ON.  WORKBOOK PAGE 8.

TIMING: 38:00-41:00. THREE MINUTES. About twenty-five seconds each.

FACILITATION: read the category and its line, move to the next. Resist elaborating.

Read each category and its one-line reading. Do not work examples, there is no room for them in sixty minutes.

THREE READINGS WERE REWRITTEN IN v2.0.6 AND THE WORDING MATTERS. Translate what is built is about making existing capability legible outside your own context, not about solving recognition. Widen exposure is about reaching people who can evaluate or use the capability, and the line says plainly that visibility is not evidence. Repair formation conditions is a category to test, not an instruction to stay.

Prepare for exit is preparation, not a recommendation to resign. Say it that way.

Say the bottom line as written: no square prescribes a category. Someone in Depth Trap is not instructed to translate; someone in Fragile is not instructed to exit."""

NOTES[22] = """RECORDING ON. PRIVATE WRITING.  WORKBOOK PAGE 8.

TIMING: 41:00-45:00. FOUR MINUTES. This is a real writing period, protect it.

FACILITATION: read the four portability questions once, read the three prompts once, then stop talking for the rest of the block.

THE FOUR-QUESTION PORTABILITY AUDIT IS ON THIS SLIDE AND IS SPOKEN ONCE, HERE. It is the interpretive frame for the note, not a new exercise: there is nothing to score and nothing to hand in. Read them and move on:
What travels? What does not? What can I prove? What must I relearn?

SAY THE DESTINATION BOUNDARY IN ONE LINE. You cannot fully answer what travels until you know where it is travelling to. Generic portability is not destination-specific portability, and this session has not named anyone's destination. What the four questions do here is tell you which of them your evidence can already answer and which one is still open.

Read the three prompts once, then stop talking and let people write. Four minutes feels long from the front and short from the seat.

Do not call it the Next-Move Decision. That is the Intensive Edition's artifact and it is a different, longer instrument.

Do not ask anyone to read theirs aloud, and do not invite anyone to share a move category, an employer or a situation.

Give a sixty-second warning at 44:00."""

NOTES[23] = """RECORDING ON.

TIMING: 45:00-46:00. SIXTY SECONDS. The continuation sequence is the first place to take time from if you are behind.

FACILITATION: read the two routes in order, land route one, then move. Do not linger and do not add examples.

ROUTE 1 IS NOT A THROWAWAY. Say it first and say it plainly. The credibility of the next slide depends on genuinely releasing the people who have enough.

WHAT TO SAY: "Some of you now have enough to act. Good, you do not need anything else from me today. Some of you want a structured system you can keep working with privately. Different needs, different routes."

TWO ROUTES ONLY IN THIS DELIVERY. Career Move Review is the current advisory offer for someone with a specific destination, and it is NOT on this slide, because its fulfillment route is not operational. Do not describe it, do not price it and do not take bookings for it from the stage. If a participant asks for a human read against a specific move, say that it is not open yet and take nothing further.

Keep the Proof is not on this slide either, and that is deliberate. It answers a different problem and belongs in follow-up, not in the recorded continuation sequence.

NO PRICING ON THIS SLIDE. Do not say "if you want to buy", "I have two offers" or "special opportunity"."""

NOTES[24] = """RECORDING ON.

TIMING: 46:00-47:00. SIXTY SECONDS.

FACILITATION: name what today already gave them, then name what the Field Kit adds. Sixty seconds. If you are behind, this is the first slide to shorten.

TODAY IS COMPLETE. Say so. This session delivers a full guided reading. The Field Kit is not the rest of it. It is a private system they can own, revisit and rerun as conditions change.

Do not call the Field Kit the complete version of the session, the full assessment, required, included, a bundle, or more worksheets.

THE CURRENT INDIVIDUAL OFFER ARCHITECTURE, for your own reference and not for the stage: the Career Evidence Starter is free, Keep the Proof is $49, the Capability Formation Field Kit is $150, and the Career Move Review is $500 with qualification before payment. Only the Field Kit appears in this recorded continuation sequence.

THE PRIVATE CAPABILITY POSITION READ IS RETIRED. It is not an offer, it is not coming back under that name, and it is not mentioned here, on slide 23 or in Q&A. Career Move Review is the current advisory architecture, and it stays off the participant-facing sequence until its route is verified operational.

IMMEDIATELY BEFORE DELIVERY, VERIFY THE CURRENT PUBLIC PRICE AND PURCHASE ROUTE. If both remain current, post the purchase link and current price in the live chat. The permanent replay slide remains price-free.

PREBUILT CHAT MESSAGE:
Continue privately with the Capability Formation Field Kit:
https://temidayoafonja.com/fieldkit

Visible branded route on the slide face: temidayoafonja.com/fieldkit. The CTA, the URL and the QR all resolve to that page, which is the stable evergreen destination, because a recorded replay outlives any storefront link. The live page remains the source of truth for current pricing.

NO PRICE APPEARS ON THIS SLIDE FACE BY DESIGN. No discount, no attendee price, no coupon, no bundle, no crossed-out price, no urgency.

The former paid group workshop is retired from active public sale and is NOT the current public next step. Do not mention it."""

NOTES[25] = """RECORDING ON — THIS IS THE LAST RECORDED SLIDE.  WORKBOOK PAGE 8 (rescore date line).

TIMING: 47:00-50:00. The close is delivered 47:00-48:00; 48:00-50:00 is FACILITATION BUFFER and this slide holds.

FACILITATION: this close returns to the question the session opened with. Say it plainly and at the pace of ordinary speech. It is not a motivational ending, there is no call to believe in yourself, and it does not need lifting.

WHAT TO SAY, close to the slide: you came in with a stay-or-leave question, I have not made that decision for you and I should not, and what you have now is a clearer read of what this work is building, what may travel, where the evidence is thin and what deserves testing next. Then the rescore date.

If you are on time, use the buffer to restate the rescore date instruction and let the room finish writing. If you are behind, the buffer is where you absorb it.

EVERGREEN-SAFE CLOSE. Do NOT say "tonight", "this cohort", "September 23", or "before enrollment closes". This has to work for someone watching the replay months from now.

AT 50:00 — STOP THE RECORDING BEFORE ADVANCING. STOP IT; DO NOT PAUSE IT. Look at the screen and confirm it has actually stopped before the first question is taken. If the platform records the whole meeting automatically, keep the raw file private and trim the distributed replay to end here.

PRE-SESSION: if the Maven/Zoom integration starts recording automatically when the host enters early, PAUSE during the pre-session holding period, then start or resume on slide 1 at 6:00 PM CT and STOP at 6:50 PM CT. Pausing belongs to the pre-session window only. The distributed recording must contain minutes 0-50 and nothing else.

Then say out loud that the recording has stopped, and only then advance.

RECORDING OFF."""

NOTES[26] = """RECORDING OFF. CONFIRM ON SCREEN THAT IT HAS STOPPED BEFORE TAKING THE FIRST QUESTION.

TIMING: 50:00-60:00. TEN MINUTES. HARD STOP AT 60:00.

FACILITATION: confirm on screen that recording has stopped, say so aloud, then take the first question. Watch the clock and close at 60:00.

This block is never distributed. It is excluded from the evergreen replay under every distribution model, including a registration-gated one. Participant questions can contain employer names, personal circumstances and third-party names, which is exactly why it sits outside the recorded core.

Q&A MAY clarify the method, the evidence protocol, boundary and incomplete readings, and move-category distinctions.

Q&A MUST NOT become individual stay/leave coaching. Do not tell anyone to resign or stay. Do not predict layoffs. Do not give legal, medical, financial, immigration or mental-health advice. Do not ask for scores, states or employer names.

If a question is really a request for a personal decision, name the boundary once and stop there.

ON DESTINATION-SPECIFIC PORTABILITY. If someone asks whether their experience will transfer to a named role, function, industry or employer, the honest answer is that this session did not read that destination. The four questions tell them what their evidence can already support and what is still open. Do not estimate their odds, and do not imply that better wording or better documentation decides what another organization will do.

THE PRIVATE CAPABILITY POSITION READ IS RETIRED AND IS NOT OFFERED. Career Move Review is the current advisory architecture and is not being sold from this stage while its route is unverified. If someone asks for a human read against a specific move, say it is not open yet. You may point to the Field Kit if it genuinely fits the question, or to nothing at all. Neither is required.

Appendix material may be used here if a question calls for it.

Close at 60:00 and thank the room."""

NOTES[33] = """PRE-SESSION HOLDING SLIDE. NOT RECORDED, EVER. It carries the only date in the deck and it is shown before the recording starts.

Up from 5:45 PM CT with the recording OFF. At 6:00 PM CT start the recording on slide 1.

OPTIONAL ANONYMOUS ARRIVAL POLL, PRE-RECORDING ONLY. If the platform makes an anonymous poll easy, it may run while this slide is up. If it does not, skip it. It is optional and nothing depends on it.

The one question is "What brought you here?" and the options are: my work feels static; I am considering an internal move; I am considering a career pivot; my role has changed around me; I want to know what would still count somewhere else.

NEVER ASK, HERE OR ANYWHERE: what state someone is in, their score, their employer, whether they plan to quit, or whether they expect layoffs. Nothing diagnostic, nothing score-revealing, nothing decision-revealing.

Close the poll before the recording starts. Do not read individual answers aloud, do not report a distribution into the recorded core, and do not refer back to it once recording begins."""


# ── helpers ─────────────────────────────────────────────────────────────────
def shapes(slide):
    return list(slide.shapes)


def set_text(shape, text):
    """Rewrite a text frame's content, keeping the first run's formatting.

    Paragraph breaks in `text` become real paragraphs cloned from the first, so
    a two-line rule stays a two-line rule with its own styling.
    """
    tf = shape.text_frame
    first = tf.paragraphs[0]
    keep = next((r for r in first.runs if r.text.strip()), first.runs[0] if first.runs
                else first.add_run())
    for p in tf.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    for r in list(first.runs):
        if r._r is not keep._r:
            r._r.getparent().remove(r._r)
    lines = text.split("\n")
    keep.text = lines[0]
    for extra in lines[1:]:
        el = copy.deepcopy(first._element)
        first._element.getparent().append(el)
        from pptx.text.text import _Paragraph
        para = _Paragraph(el, tf)
        k2 = next((r for r in para.runs if r.text.strip()), para.runs[0])
        for r in list(para.runs):
            if r._r is not k2._r:
                r._r.getparent().remove(r._r)
        k2.text = extra


def place(shape, geom):
    top, left, w, h = geom
    shape.top, shape.left, shape.width, shape.height = (
        Inches(top), Inches(left), Inches(w), Inches(h))


def build():
    os.makedirs(OUT, exist_ok=True)
    shutil.copyfile(SRC, DST)
    p = Presentation(DST)
    sl = list(p.slides)
    assert len(sl) == 33, "slide count moved"

    # ── faces ───────────────────────────────────────────────────────────────
    for idx, edits in FACE_EDITS.items():
        shp = shapes(sl[idx - 1])
        for j, expect, new in edits:
            cur = shp[j].text_frame.text
            assert expect in cur, \
                f"slide {idx} shape {j}: expected {expect[:40]!r}, found {cur[:60]!r}"
            set_text(shp[j], new)

    # ── slide 18: drop the printed procedure, then re-lay out ───────────────
    s18 = sl[17]
    shp = shapes(s18)
    for j in S18_DELETE:
        shp[j]._element.getparent().remove(shp[j]._element)
    for j, geom in S18_GEOM.items():
        place(shp[j], geom)
    # The repurposed line becomes a rule, so it is set in bold rather than left
    # reading as one more body sentence.
    for para in shp[16].text_frame.paragraphs:
        for r in para.runs:
            r.font.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = NAVY

    # Pre-existing defect, repaired here: the third label ran past its 5.40in
    # box and collided with the answer rule at 6.10in. Shortened to the length
    # the other two labels use.
    s22_label = shapes(sl[21])[18]
    assert "One action, conversation, or piece of evidence" in s22_label.text_frame.text
    set_text(s22_label, "The next action, conversation or piece of evidence is:")

    # ── slide 22: the portability frame, as a second line of the subtitle ───
    s22 = sl[21]
    shp22 = shapes(s22)
    place(shp22[2], S22_GEOM[2])
    tf = shp22[2].text_frame
    assert "Three lines." in tf.text, "slide 22 subtitle moved"
    assert "What travels?" not in tf.text, "slide 22 already carries the frame"
    first = tf.paragraphs[0]
    el = copy.deepcopy(first._element)
    first._element.getparent().append(el)
    from pptx.text.text import _Paragraph
    para = _Paragraph(el, tf)
    keep = next(r for r in para.runs if r.text.strip())
    for r in list(para.runs):
        if r._r is not keep._r:
            r._r.getparent().remove(r._r)
    keep.text = PORTABILITY_FRAME
    keep.font.bold = True
    keep.font.size = Pt(10.0)
    keep.font.color.rgb = NAVY

    # ── slide 25: the close needs another line of height ────────────────────
    place(shapes(sl[24])[3], S25_GEOM[3])

    # ── notes ───────────────────────────────────────────────────────────────
    for idx, text in NOTES.items():
        sl[idx - 1].notes_slide.notes_text_frame.text = text

    p.save(DST)
    return DST


if __name__ == "__main__":
    path = build()
    print("built", os.path.basename(path))
    print("  sha256", hashlib.sha256(open(path, "rb").read()).hexdigest())
